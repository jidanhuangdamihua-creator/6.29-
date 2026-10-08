#!/usr/bin/env python3
"""180-day overlay for the 3 paper formal targets.

Recalculates SITA + 8 baselines under additional-target / 70-target calendars
(n_test=180, every test day is a label). Does not reuse 170/168/166 paper1target
cells, does not rewrite 60-target manifests, additional_target_robustness
internals, official source_model_cache, old paper1target xlsx, or fixed40_headline.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import shutil
import sys
import traceback
from datetime import datetime, timezone
from pathlib import Path

os.environ.setdefault("KMP_DUPLICATE_LIB_OK", "TRUE")
os.environ.setdefault("OMP_NUM_THREADS", "1")
os.environ.setdefault("CUDA_VISIBLE_DEVICES", "")
os.environ.setdefault("TF_CPP_MIN_LOG_LEVEL", "2")

import numpy as np
import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from additional_target_robustness.horizon import verify_direct_horizon
from additional_target_robustness.paths import SEALED_ROOT, SOURCE_MODEL_CACHE, ensure_ablation_path
from additional_target_robustness.protocol import (
    HORIZONS,
    PAPER_TARGETS,
    SEEDS,
    TEST_DAYS,
    TRAIN_DAYS,
    VAL_DAYS,
    Calendar,
    ProtocolError,
    build_evaluation_index,
    calendar_for,
    frozen_protocol_sha256,
    sha256_file,
)
from additional_target_robustness.runner import (
    _load_or_train_source,
    _one_source,
    _predict_sita,
    _series_sha256,
)
from additional_target_robustness.source_selection import load_source_histories, select_top3
from additional_target_robustness.target_data import (
    _columns,
    _read_filtered,
    sales_by_date,
    split_target_frame,
)
import additional_target_robustness.runner as sita_runner

ensure_ablation_path()
from sita_ablation.metrics import evaluate_original_sales

OUT_ROOT = REPO_ROOT / "outputs" / "paper3target_180day_20261005"
PLAN_DIR = REPO_ROOT / "实验结果" / "paper3target_180day_plan"
XLSX_DIR = REPO_ROOT / "实验结果"
COHORT = "PAPER_FORMAL_1_180"
COHORT_LABEL = "paper formal 1-target restated on 180-day robustness protocol"
SELECTION_ORIGIN = "PAPER_FORMAL_1_180"
FORBIDDEN_D4 = {"166/311", "166-311", "166_311"}
METHODS = (
    "SITA",
    "Seasonal Naive",
    "LightGBM",
    "LSTM",
    "DLinear",
    "TCN",
    "PatchTST",
    "TimesNet",
    "MSML-TL",
)
BL_METHODS = METHODS[1:]
RECYCLE_AFTER = 15
DATASETS = {
    "D4": {
        "identity": "166/258",
        "store": "166",
        "item": "258",
        "source_path": "数据集/固化数据/d1_d6_sealed_v1/dataset4/source.parquet",
        "xlsx": "D4_paper1target_BL_SITA_180.xlsx",
        "old_xlsx": "D4_paper1target_BL_SITA.xlsx",
        "workbook_70": "D4_70target_BL_SITA.xlsx",
        "group_field": "same_store_only",
        "group_value": "True",
        "bl_module": "d4_multistore_baselines",
    },
    "D5": {
        "identity": "48/1159414",
        "store": "48",
        "item": "1159414",
        "source_path": "数据集/固化数据/d1_d6_sealed_v1/dataset5/prepared_180day_source.parquet",
        "xlsx": "D5_paper1target_BL_SITA_180.xlsx",
        "old_xlsx": "D5_paper1target_BL_SITA.xlsx",
        "workbook_70": "D5_70target_BL_SITA.xlsx",
        "group_field": "family",
        "group_value": "GROCERY I",
        "bl_module": "d5_d6_60target_baselines",
    },
    "D6": {
        "identity": "CA_1/FOODS_3_080",
        "store": "CA_1",
        "item": "FOODS_3_080",
        "source_path": "数据集/固化数据/d1_d6_sealed_v1/dataset6/source.parquet",
        "xlsx": "D6_paper1target_BL_SITA_180.xlsx",
        "old_xlsx": "D6_paper1target_BL_SITA.xlsx",
        "workbook_70": "D6_70target_BL_SITA.xlsx",
        "group_field": "dept_id",
        "group_value": "FOODS_3",
        "bl_module": "d5_d6_60target_baselines",
    },
}


class OverlayStop(RuntimeError):
    def __init__(self, token: str, detail: str = "") -> None:
        super().__init__(token if not detail else f"{token}: {detail}")
        self.token = token
        self.detail = detail


class _Log:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)

    def write(self, message: str) -> None:
        stamp = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        line = f"[{stamp}] {message}"
        print(line, flush=True)
        with self.path.open("a", encoding="utf-8") as handle:
            handle.write(line + "\n")


def _utc() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _atomic_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    tmp.replace(path)


def _atomic_csv(path: Path, frame: pd.DataFrame) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    frame.to_csv(tmp, index=False)
    tmp.replace(path)


def _payload_digest(payload: dict) -> str:
    return hashlib.sha256(json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()


def _cfg(dataset: str) -> dict:
    return DATASETS[dataset]


def _paper_target_path(dataset: str) -> Path:
    folder = {"D4": "dataset4", "D5": "dataset5", "D6": "dataset6"}[dataset]
    path = SEALED_ROOT / folder / "target.parquet"
    if not path.is_file():
        raise OverlayStop("HOLD_TARGET_MANIFEST_MISMATCH", str(path))
    return path


def load_paper_target_series(dataset: str, identity: str) -> pd.DataFrame:
    """Read the paper identity from sealed target.parquet, not source.parquet.

    additional_target_robustness.load_target_series points at source.parquet for
    sampled 70-target identities. The three paper keys were excluded from that
    sampling universe and only exist on the target authority.
    """
    calendar = calendar_for(dataset)
    path = _paper_target_path(dataset)
    date_max = calendar.test_end if dataset != "D5" else "2017-08-14"
    frame = _read_filtered(
        path,
        key_fields=calendar.key_fields,
        identity=identity,
        date_min=calendar.train_start,
        date_max=date_max,
        columns=_columns(calendar),
    )
    frame["date"] = pd.to_datetime(frame["date"]).dt.normalize()
    if dataset == "D5":
        cutoff = pd.Timestamp("2017-08-14")
        banned = pd.Timestamp("2017-08-15")
        frame = frame.loc[frame["date"] <= cutoff].copy()
        if (frame["date"] == banned).any() or (frame["date"] > cutoff).any():
            raise OverlayStop("BLOCKED_D5_CUTOFF", f"{identity} still contains 2017-08-15")
    if frame.empty:
        raise OverlayStop("HOLD_TARGET_MANIFEST_MISMATCH", f"{dataset} {identity} missing from target.parquet")
    train, validation, test = split_target_frame(frame, calendar)
    del train, validation, test
    return frame.sort_values("date", kind="mergesort").reset_index(drop=True)


def _identity(dataset: str) -> str:
    identity = _cfg(dataset)["identity"]
    if dataset == "D4" and identity in FORBIDDEN_D4:
        raise OverlayStop("BLOCKED_D4_TARGET", "refusing 166/311")
    if PAPER_TARGETS[dataset] != identity:
        raise OverlayStop("BLOCKED_PAPER_TARGET", f"{dataset} {identity} != {PAPER_TARGETS[dataset]}")
    return identity


def _ds_root(dataset: str) -> Path:
    return OUT_ROOT / dataset


def _overlay_cache() -> Path:
    return OUT_ROOT / "source_model_cache"


def _selection_path(dataset: str) -> Path:
    return _ds_root(dataset) / "selection.json"


def _targets_meta_path(dataset: str) -> Path:
    return _ds_root(dataset) / "targets.json"


def _sita_cell_dir(dataset: str, identity: str, horizon: int, seed: int) -> Path:
    safe = identity.replace("/", "_")
    return _ds_root(dataset) / "sita" / "cells" / f"{safe}__h{horizon}__s{seed}"


def _load_sheet(path: Path, name: str) -> tuple[list, list[dict]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb[name].iter_rows(values_only=True))
    wb.close()
    hdr = list(rows[0])
    return hdr, [dict(zip(hdr, row)) for row in rows[1:]]


def _in_70target(dataset: str, identity: str) -> bool:
    path = XLSX_DIR / _cfg(dataset)["workbook_70"]
    if not path.is_file():
        return False
    _, rows = _load_sheet(path, "targets")
    return any(str(row.get("target_identity")) == identity for row in rows)


def _audit_old_xlsx(dataset: str) -> dict:
    path = XLSX_DIR / _cfg(dataset)["old_xlsx"]
    calendar = calendar_for(dataset)
    if not path.is_file():
        return {"present": False, "path": str(path.relative_to(REPO_ROOT))}
    _, rows = _load_sheet(path, "cell_raw")
    n_tests = sorted({int(row["n_test"]) for row in rows if row.get("n_test") is not None})
    by_h = {}
    for row in rows:
        key = (str(row["method"]), int(row["horizon"]))
        by_h.setdefault(key, {"n_test": set(), "first": set(), "last": set()})
        if row.get("n_test") is not None:
            by_h[key]["n_test"].add(int(row["n_test"]))
        if row.get("first_target_date"):
            by_h[key]["first"].add(str(row["first_target_date"])[:10])
        if row.get("last_target_date"):
            by_h[key]["last"].add(str(row["last_target_date"])[:10])
    compact = {
        f"{method}|h{horizon}": {
            "n_test": sorted(vals["n_test"]),
            "first_target_date": sorted(vals["first"]),
            "last_target_date": sorted(vals["last"]),
        }
        for (method, horizon), vals in sorted(by_h.items())
    }
    d5_has_aug15 = False
    if dataset == "D5":
        for row in rows:
            for field in ("first_target_date", "last_target_date"):
                text = str(row.get(field) or "")
                if "2017-08-15" in text:
                    d5_has_aug15 = True
    return {
        "present": True,
        "path": str(path.relative_to(REPO_ROOT)),
        "n_rows": len(rows),
        "n_test_values": n_tests,
        "uses_170_168_166": set(n_tests) <= {166, 167, 168, 169, 170, 171} and 180 not in n_tests,
        "calendar_test_start": calendar.test_start,
        "calendar_test_end": calendar.test_end,
        "d5_has_2017_08_15": d5_has_aug15,
        "by_method_horizon": compact,
        "reusable_as_180": False,
        "reuse_block_reason": "n_test is 170/168/166 (or D5 MSML 171/169/167); scoring support is not every test day",
    }


def _expected_index(dataset: str) -> dict:
    calendar = calendar_for(dataset)
    payload = {
        "dataset": dataset,
        "train_start": calendar.train_start,
        "train_end": calendar.train_end,
        "val_start": calendar.val_start,
        "val_end": calendar.val_end,
        "test_start": calendar.test_start,
        "test_end": calendar.test_end,
        "n_train": TRAIN_DAYS,
        "n_val": VAL_DAYS,
        "n_test": TEST_DAYS,
        "horizons": {},
    }
    for horizon in HORIZONS:
        index = build_evaluation_index(dataset, horizon)
        dates = index.label_date_text()
        payload["horizons"][str(horizon)] = {
            "n_test": len(dates),
            "first_target_date": dates[0],
            "last_target_date": dates[-1],
            "matches_test_start": dates[0] == calendar.test_start,
            "matches_test_end": dates[-1] == calendar.test_end,
        }
        if dataset == "D5" and (dates[-1] > "2017-08-14" or "2017-08-15" in dates):
            raise OverlayStop("BLOCKED_D5_CUTOFF", f"h={horizon} last={dates[-1]}")
    return payload


def phase_audit() -> dict:
    PLAN_DIR.mkdir(parents=True, exist_ok=True)
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    contract = verify_direct_horizon()
    datasets = {}
    for dataset in DATASETS:
        identity = _identity(dataset)
        calendar = calendar_for(dataset)
        old = _audit_old_xlsx(dataset)
        in70 = _in_70target(dataset, identity)
        existing_180 = list((_ds_root(dataset) / "sita" / "cells").glob("*/DONE.json")) if _ds_root(dataset).exists() else []
        datasets[dataset] = {
            "identity": identity,
            "in_70target_workbook": in70,
            "old_paper1target": old,
            "expected": _expected_index(dataset),
            "existing_overlay_sita_done": len(existing_180),
            "official_60cell_cache_reusable": False,
            "d5_cutoff_block": dataset == "D5" and bool(old.get("d5_has_2017_08_15")),
        }
        if dataset == "D4":
            datasets[dataset]["forbidden_166_311"] = False
    payload = {
        "status": "AUDIT_OK",
        "cohort": COHORT,
        "contract": contract,
        "protocol_sha256": frozen_protocol_sha256(),
        "official_source_model_cache": str(SOURCE_MODEL_CACHE.relative_to(REPO_ROOT)),
        "overlay_root": str(OUT_ROOT.relative_to(REPO_ROOT)),
        "do_not_reuse": [
            "实验结果/D*_paper1target_BL_SITA.xlsx (170/168/166)",
            "outputs/paper3target_sita_isolated_20261005",
            "outputs/paper3target_8bl_isolated_20261005",
            "corrected 60-cell SITA cells",
        ],
        "do_not_rewrite": [
            "additional_target_robustness internals",
            "D*_70target_BL_SITA.xlsx",
            "D*_paper1target_BL_SITA.xlsx",
            "outputs/additional_target_robustness_20261004/source_model_cache",
            "实验结果/sita_d456_fixed40_headline/",
        ],
        "datasets": datasets,
        "audited_at": _utc(),
    }
    _atomic_json(PLAN_DIR / "00_audit.json", payload)
    lines = [
        "# paper 3-target 180-day overlay — Phase A audit",
        "",
        f"Written at `{payload['audited_at']}`.",
        "",
        "## Verdict",
        "",
        "- Do **not** merge current `*_paper1target_BL_SITA.xlsx` into fixed40.",
        "- All three identities are **absent** from the 70-target workbooks (explicitly excluded as formal paper targets).",
        "- Official 60-cell / paper1target caches are **not reusable**; the label set changed.",
        "- D5 `2017-08-15` in the old paper1target table is a hard block for 180-day restatement.",
        "",
        "## Calendar (locked)",
        "",
        "| dataset | identity | train | val | test | n_test |",
        "|---|---|---|---|---|---|",
    ]
    for dataset, item in datasets.items():
        exp = item["expected"]
        lines.append(
            f"| {dataset} | `{item['identity']}` | {exp['train_start']}–{exp['train_end']} | "
            f"{exp['val_start']}–{exp['val_end']} | {exp['test_start']}–{exp['test_end']} | {exp['n_test']} |"
        )
    lines.extend(
        [
            "",
            "## Old paper1target vs 180-day",
            "",
            "| dataset | old n_test | old first/last (h1 SITA if present) | 180 first | 180 last | reusable |",
            "|---|---|---|---|---|---|",
        ]
    )
    for dataset, item in datasets.items():
        old = item["old_paper1target"]
        exp = item["expected"]["horizons"]["1"]
        h1 = (old.get("by_method_horizon") or {}).get("SITA|h1", {})
        lines.append(
            f"| {dataset} | {old.get('n_test_values')} | {h1.get('first_target_date')}/{h1.get('last_target_date')} | "
            f"{exp['first_target_date']} | {exp['last_target_date']} | **no** |"
        )
    lines.extend(
        [
            "",
            "## Suggested command order",
            "",
            "```bash",
            "python scripts/paper3target_180day.py --phase prepare --dataset D4",
            "python scripts/paper3target_180day.py --phase prepare --dataset D5",
            "python scripts/paper3target_180day.py --phase prepare --dataset D6",
            "python scripts/paper3target_180day.py --phase stage-a --dataset D4",
            "python scripts/paper3target_180day.py --phase stage-a --dataset D5",
            "python scripts/paper3target_180day.py --phase stage-a --dataset D6",
            "python scripts/paper3target_180day.py --phase sita --dataset D4",
            "python scripts/paper3target_180day.py --phase sita --dataset D5",
            "python scripts/paper3target_180day.py --phase sita --dataset D6",
            "python scripts/paper3target_180day.py --phase baselines --dataset D4",
            "python scripts/paper3target_180day.py --phase baselines --dataset D5",
            "python scripts/paper3target_180day.py --phase baselines --dataset D6",
            "python scripts/paper3target_180day.py --phase workbook",
            "```",
            "",
            "Do not start `--phase headline` until every cell has `n_test=180`.",
            "",
        ]
    )
    (PLAN_DIR / "00_audit.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return payload


def _bind_overlay_cache() -> Path:
    cache = _overlay_cache()
    cache.mkdir(parents=True, exist_ok=True)
    sita_runner.SOURCE_MODEL_CACHE = cache
    return cache


def _stage_a_payload(dataset: str, source_identity: str, horizon: int, seed: int, series: pd.DataFrame, features) -> dict:
    return {
        "dataset": dataset,
        "source_identity": source_identity,
        "horizon": int(horizon),
        "seed": int(seed),
        "architecture": "cnn",
        "source_data_sha256": _series_sha256(series, features),
        "protocol_sha256": frozen_protocol_sha256(),
    }


def _copy_official_hit(digest: str) -> str:
    src = SOURCE_MODEL_CACHE / digest
    dst = _overlay_cache() / digest
    if (dst / "DONE.json").is_file() and (dst / "weights.npz").is_file():
        return "OVERLAY_HIT"
    if (src / "DONE.json").is_file() and (src / "weights.npz").is_file():
        if dst.exists():
            shutil.rmtree(dst)
        shutil.copytree(src, dst)
        return "COPIED_OFFICIAL_HIT"
    return "MISS"


def phase_prepare(dataset: str) -> dict:
    identity = _identity(dataset)
    log = _Log(_ds_root(dataset) / "prepare_log.txt")
    log.write(f"prepare start {dataset} {identity}")
    calendar = calendar_for(dataset)
    target = load_paper_target_series(dataset, identity)
    train, validation, test = split_target_frame(target, calendar)
    if dataset == "D5":
        cutoff = pd.Timestamp("2017-08-14")
        if (target["date"] > cutoff).any() or (target["date"] == pd.Timestamp("2017-08-15")).any():
            raise OverlayStop("BLOCKED_D5_CUTOFF", identity)
    if len(train) != TRAIN_DAYS or len(validation) != VAL_DAYS or len(test) != TEST_DAYS:
        raise OverlayStop("BLOCKED_SITA_PROTOCOL", f"{dataset} split lengths")
    histories = load_source_histories(dataset, target_identity=identity, target_frame=target)
    universe = f"{COHORT}|{dataset}|store={identity.split('/')[0]}"
    selection = select_top3(dataset, identity, target, histories, source_universe_id=universe)
    if selection.status != "COMPLETED" or len(selection.selected) != 3:
        raise OverlayStop("HOLD_SOURCE_SELECTION_DRIFT", f"{dataset} {identity} status={selection.status}")
    live = [item.identity for item in selection.selected]
    if identity in live:
        raise OverlayStop("HOLD_TARGET_NOT_EXCLUDED", identity)
    store = identity.split("/")[0]
    if any(item.split("/")[0] != store for item in live):
        raise OverlayStop("HOLD_SOURCE_SELECTION_DRIFT", f"cross-store {live}")
    features = calendar.model_features
    key_fields = calendar.key_fields
    requirements = []
    for source in selection.selected:
        series = _one_source(histories, source.key, key_fields)
        for horizon in HORIZONS:
            for seed in SEEDS:
                payload = _stage_a_payload(dataset, source.identity, horizon, seed, series, features)
                digest = _payload_digest(payload)
                status = _copy_official_hit(digest)
                requirements.append(
                    {
                        "dataset": dataset,
                        "source_identity": source.identity,
                        "horizon": int(horizon),
                        "seed": int(seed),
                        "cache_key": digest,
                        "status": status,
                    }
                )
    n_miss = sum(1 for row in requirements if row["status"] == "MISS")
    n_hit = len(requirements) - n_miss
    selection_payload = {
        "cohort": COHORT,
        "dataset": dataset,
        "target_identity": identity,
        "status": selection.status,
        "legal_count": selection.legal_count,
        "comparison_start": selection.comparison_start,
        "comparison_end": selection.comparison_end,
        "source_universe_id": universe,
        "selected": [
            {
                "rank": int(item.rank) + 1,
                "identity": item.identity,
                "distance": float(item.distance),
                "weight": float(item.weight),
            }
            for item in selection.selected
        ],
        "prepared_at": _utc(),
    }
    targets_payload = {
        "cohort": COHORT,
        "dataset": dataset,
        "target_identity": identity,
        "target_rank": 1,
        "selection_origin": SELECTION_ORIGIN,
        "eligibility_status": "paper_formal_target_180",
        "train_start": calendar.train_start,
        "train_end": calendar.train_end,
        "val_start": calendar.val_start,
        "val_end": calendar.val_end,
        "test_start": calendar.test_start,
        "test_end": calendar.test_end,
        "n_train_days": TRAIN_DAYS,
        "n_val_days": VAL_DAYS,
        "n_test_days": TEST_DAYS,
        "source_universe_id": universe,
        "source_data_path": _cfg(dataset)["source_path"],
        "protocol_version": "paper_formal_1_180",
        "n_rows": int(len(target)),
        "min_date": pd.Timestamp(target["date"].min()).strftime("%Y-%m-%d"),
        "max_date": pd.Timestamp(target["date"].max()).strftime("%Y-%m-%d"),
    }
    req_frame = pd.DataFrame(requirements)
    _atomic_json(_selection_path(dataset), selection_payload)
    _atomic_json(_targets_meta_path(dataset), targets_payload)
    _atomic_csv(_ds_root(dataset) / "stage_a_requirements.csv", req_frame)
    summary = {
        "status": "PREPARE_OK" if n_miss == 0 else "PREPARE_NEEDS_STAGE_A",
        "dataset": dataset,
        "target_identity": identity,
        "selected": live,
        "n_stage_a_required": len(requirements),
        "n_stage_a_hit": n_hit,
        "n_stage_a_miss": n_miss,
        "test_start": calendar.test_start,
        "test_end": calendar.test_end,
        "finished_at": _utc(),
    }
    _atomic_json(_ds_root(dataset) / "prepare_summary.json", summary)
    log.write(json.dumps(summary, sort_keys=True))
    return summary


def _load_prepared(dataset: str):
    identity = _identity(dataset)
    selection = json.loads(_selection_path(dataset).read_text(encoding="utf-8"))
    calendar = calendar_for(dataset)
    target = load_paper_target_series(dataset, identity)
    train, validation, test = split_target_frame(target, calendar)
    histories = load_source_histories(dataset, target_identity=identity, target_frame=target)
    universe = selection["source_universe_id"]
    live = select_top3(dataset, identity, target, histories, source_universe_id=universe)
    expected = [item["identity"] for item in selection["selected"]]
    got = [item.identity for item in live.selected]
    if got != expected:
        raise OverlayStop("HOLD_SOURCE_SELECTION_DRIFT", f"live={got} frozen={expected}")
    return {
        "identity": identity,
        "calendar": calendar,
        "target": target,
        "train": train,
        "validation": validation,
        "test": test,
        "histories": histories,
        "selection": live,
        "selected": got,
        "sales": sales_by_date(target),
        "meta": selection,
    }


def _hit_only_source(histories, *, source_key, features, horizon, seed, dataset, key_fields):
    series = _one_source(histories, source_key, key_fields)
    identity = "/".join(source_key)
    payload = _stage_a_payload(dataset, identity, horizon, seed, series, features)
    digest = _payload_digest(payload)
    directory = _overlay_cache() / digest
    done_path = directory / "DONE.json"
    if not done_path.is_file():
        raise OverlayStop("HOLD_STAGE_A_CACHE_MISS", f"{dataset} {identity} h={horizon} seed={seed} key={digest}")
    return sita_runner._read_source_record(directory, source_key), "CACHE_HIT"


def _recycle(phase: str, dataset: str, log: _Log, newly: int, leftover: int) -> None:
    if newly < RECYCLE_AFTER or leftover <= 0:
        return
    log.write(f"recycling TensorFlow after {newly} cells leftover={leftover} phase={phase} dataset={dataset}")
    os.execv(sys.executable, [sys.executable, str(Path(__file__).resolve()), "--phase", phase, "--dataset", dataset])


def phase_stage_a(dataset: str) -> dict:
    identity = _identity(dataset)
    log = _Log(_ds_root(dataset) / "stage_a_log.txt")
    req_path = _ds_root(dataset) / "stage_a_requirements.csv"
    if not req_path.is_file():
        raise OverlayStop("HOLD_STAGE_A_CACHE_MISS", "run --phase prepare first")
    _bind_overlay_cache()
    prepared = _load_prepared(dataset)
    req = pd.read_csv(req_path)
    calendar = prepared["calendar"]
    newly = 0
    pending = req.loc[req["status"] == "MISS"]
    leftover = int(len(pending))
    for _, row in pending.iterrows():
        source_identity = str(row["source_identity"])
        horizon = int(row["horizon"])
        seed = int(row["seed"])
        digest = str(row["cache_key"])
        copied = _copy_official_hit(digest)
        if copied != "MISS":
            continue
        log.write(f"{dataset} stage-a TRAIN {source_identity} h={horizon} seed={seed}")
        source_key = tuple(source_identity.split("/"))
        _load_or_train_source(
            prepared["histories"],
            source_key=source_key,
            features=calendar.model_features,
            horizon=horizon,
            seed=seed,
            dataset=dataset,
            key_fields=calendar.key_fields,
        )
        newly += 1
        leftover -= 1
        _recycle("stage-a", dataset, log, newly, leftover)
    # refresh statuses
    statuses = []
    for _, row in req.iterrows():
        status = _copy_official_hit(str(row["cache_key"]))
        statuses.append(status if status != "COPIED_OFFICIAL_HIT" else "OVERLAY_HIT")
    req = req.copy()
    req["status"] = statuses
    _atomic_csv(req_path, req)
    n_miss = int((req["status"] == "MISS").sum())
    summary = {
        "status": "STAGE_A_OK" if n_miss == 0 else "STAGE_A_INCOMPLETE",
        "dataset": dataset,
        "target_identity": identity,
        "n_miss": n_miss,
        "n_hit": int(len(req) - n_miss),
        "finished_at": _utc(),
    }
    _atomic_json(_ds_root(dataset) / "stage_a_summary.json", summary)
    log.write(json.dumps(summary, sort_keys=True))
    if n_miss:
        raise OverlayStop("HOLD_STAGE_A_CACHE_MISS", f"{dataset} still missing {n_miss} keys")
    return summary


def _write_sita_cell(dataset, directory, *, cell_key, identity, horizon, seed, dates, y_true, y_pred, selected, usage, protocol_sha):
    if len(dates) != TEST_DAYS or y_true.shape != (TEST_DAYS,) or y_pred.shape != (TEST_DAYS,):
        raise OverlayStop("BLOCKED_SITA_PROTOCOL", f"n_test {len(dates)}")
    pair = evaluate_original_sales(y_true, y_pred)
    if directory.exists():
        shutil.rmtree(directory)
    tmp = directory.with_name(directory.name + ".partial")
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir(parents=True)
    pd.DataFrame(
        {
            "dataset": dataset,
            "target_identity": identity,
            "method": "SITA",
            "horizon": int(horizon),
            "seed": int(seed),
            "target_date": list(dates),
            "y_true": y_true,
            "y_pred": y_pred,
        }
    ).to_csv(tmp / "predictions.csv", index=False)
    record = {
        "cell_key": cell_key,
        "status": "COMPLETED",
        "dataset": dataset,
        "method": "SITA",
        "target_identity": identity,
        "horizon": int(horizon),
        "seed": int(seed),
        "n_test": TEST_DAYS,
        "rmse": pair.rmse,
        "smape": pair.smape,
        "prediction_finite_count": int(np.isfinite(y_pred).sum()),
        "target_finite_count": int(np.isfinite(y_true).sum()),
        "first_target_date": dates[0],
        "last_target_date": dates[-1],
        "selected_sources": ",".join(selected),
        "cache_status": "OVERLAY_180",
        "new_source_training": False,
        "canary_cell": False,
        "protocol_sha256": protocol_sha,
        "trained_at": _utc(),
        "cohort": COHORT,
    }
    (tmp / "DONE.json").write_text(json.dumps(record, indent=2, sort_keys=True), encoding="utf-8")
    (tmp / "source_usage.json").write_text(json.dumps(usage, indent=2, sort_keys=True), encoding="utf-8")
    tmp.replace(directory)
    return record


def phase_sita(dataset: str) -> dict:
    identity = _identity(dataset)
    log = _Log(_ds_root(dataset) / "sita" / "run_log.txt")
    _bind_overlay_cache()
    sita_runner._load_or_train_source = _hit_only_source
    prepared = _load_prepared(dataset)
    calendar = prepared["calendar"]
    protocol_sha = frozen_protocol_sha256()
    selected = prepared["selected"]
    progress_path = _ds_root(dataset) / "sita" / "sita_progress.csv"
    progress = []
    if progress_path.is_file():
        progress = pd.read_csv(progress_path).to_dict("records")
    done = {
        (int(row["horizon"]), int(row["seed"]))
        for row in progress
        if str(row.get("status")) in {"COMPLETED", "SKIP"}
    }
    newly = 0
    leftover = 15 - len(done)
    for horizon in HORIZONS:
        index = build_evaluation_index(dataset, horizon)
        dates = index.label_date_text()
        if len(dates) != TEST_DAYS or dates[0] != calendar.test_start or dates[-1] != calendar.test_end:
            raise OverlayStop("BLOCKED_SITA_PROTOCOL", f"{dataset} h={horizon} index")
        y_true = np.asarray([prepared["sales"][day] for day in dates], dtype=np.float64)
        for seed in SEEDS:
            if (int(horizon), int(seed)) in done:
                continue
            directory = _sita_cell_dir(dataset, identity, horizon, seed)
            cell_key = _payload_digest(
                {
                    "cohort": COHORT,
                    "dataset": dataset,
                    "target": identity,
                    "horizon": int(horizon),
                    "seed": int(seed),
                    "n_test": TEST_DAYS,
                    "protocol_sha256": protocol_sha,
                    "sources": selected,
                }
            )
            log.write(f"{dataset} SITA {identity} h={horizon} seed={seed}")
            y_pred, usage = _predict_sita(
                prepared["target"],
                prepared["train"],
                prepared["validation"],
                prepared["selection"],
                index,
                horizon=int(horizon),
                seed=int(seed),
                dataset=dataset,
                identity=identity,
                histories=prepared["histories"],
            )
            record = _write_sita_cell(
                dataset,
                directory,
                cell_key=cell_key,
                identity=identity,
                horizon=horizon,
                seed=seed,
                dates=dates,
                y_true=y_true,
                y_pred=np.asarray(y_pred, dtype=np.float64),
                selected=selected,
                usage=usage,
                protocol_sha=protocol_sha,
            )
            progress.append(record)
            done.add((int(horizon), int(seed)))
            newly += 1
            leftover = 15 - len(done)
            _atomic_csv(progress_path, pd.DataFrame(progress))
            log.write(
                f"{dataset} SITA {identity} h={horizon} seed={seed} n_test=180 "
                f"rmse={record['rmse']:.6g} first={dates[0]} last={dates[-1]}"
            )
            _recycle("sita", dataset, log, newly, leftover)
    n_ok = sum(1 for row in progress if str(row.get("status")) in {"COMPLETED", "SKIP"})
    summary = {"status": "SITA_OK" if n_ok == 15 else "SITA_INCOMPLETE", "dataset": dataset, "n_completed": n_ok}
    _atomic_json(_ds_root(dataset) / "sita" / "summary.json", summary)
    return summary


def _load_bl(name: str):
    path = REPO_ROOT / "scripts" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise OverlayStop("BLOCKED_BASELINE_PROTOCOL", str(path))
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _patch_baselines(dataset: str):
    identity = _identity(dataset)
    overlay_hash = sha256_file(_selection_path(dataset)) if _selection_path(dataset).is_file() else COHORT
    if dataset == "D4":
        bl = _load_bl("d4_multistore_baselines")
        out = _ds_root(dataset)
        bl.OUT_DIR = out
        bl.BASELINE_ROOT = out / "baselines"
        bl.MANIFEST_PATH = _selection_path(dataset)
        bl.MANIFEST_SHA256 = overlay_hash
        bl.BASELINE_ROOT.mkdir(parents=True, exist_ok=True)

        def _recycle(log, reason: str) -> None:
            log.write(f"recycling TensorFlow ({reason})")
            os.execv(
                sys.executable,
                [sys.executable, str(Path(__file__).resolve()), "--phase", "baselines", "--dataset", "D4"],
            )

        bl._recycle_baseline = _recycle
        bl.load_target_series = load_paper_target_series
        return bl, [identity]
    bl = _load_bl("d5_d6_60target_baselines")
    bl.DATASETS[dataset]["out"] = _ds_root(dataset)
    bl.DATASETS[dataset]["manifest_sha256"] = overlay_hash
    bl.DATASETS[dataset]["manifest_name"] = "selection.json"
    bl.load_target_series = load_paper_target_series
    (_ds_root(dataset) / "baselines").mkdir(parents=True, exist_ok=True)
    return bl, [identity]


def phase_baselines(dataset: str, methods: tuple[str, ...] | None = None) -> dict:
    identity = _identity(dataset)
    selected = tuple(methods) if methods else BL_METHODS
    bl, identities = _patch_baselines(dataset)
    log = bl._Log(_ds_root(dataset) / "baselines" / "baseline_run_log.txt")
    log.write(f"{dataset} overlay baselines {identity} methods={list(selected)}")
    if dataset == "D4":
        progress = bl.run_cells(
            identities=identities,
            methods=selected,
            horizons=HORIZONS,
            seeds=SEEDS,
            log=log,
            msml_blocked=False,
        )
    else:
        progress = bl.run_cells(
            dataset,
            identities=identities,
            methods=selected,
            horizons=HORIZONS,
            seeds=SEEDS,
            log=log,
            msml_blocked=False,
        )
    n_ok = sum(1 for row in progress if str(row.get("status")) in {"COMPLETED", "SKIP"} and int(row.get("n_test") or 0) == 180)
    summary = {
        "status": "BASELINES_PARTIAL",
        "dataset": dataset,
        "n_progress_ok_180": n_ok,
        "n_progress_rows": len(progress),
        "finished_at": _utc(),
    }
    _atomic_json(_ds_root(dataset) / "baselines" / "overlay_summary.json", summary)
    return summary


def _iter_done(dataset: str) -> list[dict]:
    rows = []
    identity = _identity(dataset)
    sita_root = _ds_root(dataset) / "sita" / "cells"
    if sita_root.is_dir():
        for path in sorted(sita_root.glob("*/DONE.json")):
            rows.append(json.loads(path.read_text(encoding="utf-8")) | {"_done": str(path.relative_to(REPO_ROOT))})
    bl_root = _ds_root(dataset) / "baselines"
    slugs = {
        "Seasonal Naive": "seasonal_naive",
        "LightGBM": "lightgbm",
        "LSTM": "lstm",
        "DLinear": "dlinear",
        "TCN": "tcn",
        "PatchTST": "patchtst",
        "TimesNet": "timesnet",
        "MSML-TL": "msml_tl",
    }
    for method, slug in slugs.items():
        cells = bl_root / slug / "cells"
        if not cells.is_dir():
            continue
        for path in sorted(cells.glob("*/DONE.json")):
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["_done"] = str(path.relative_to(REPO_ROOT))
            payload.setdefault("method", method)
            rows.append(payload)
    return rows


def phase_workbook() -> list[Path]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from tools.build_d5_70target_workbook import CELL_COLUMNS

    seed_cols = [
        "cohort", "method", "store_id", "target_identity", "target_rank",
        "selection_origin", "horizon", "n_seeds", "rmse_mean", "rmse_std",
        "smape_mean", "smape_std", "n_test", "first_target_date", "last_target_date",
        "selected_sources", "deterministic_reuse",
    ]
    cov_cols = ["cohort", "method", "n_targets_expected", "n_cells_expected", "n_cells_completed", "n_cells_other", "complete"]
    tgt_cols = [
        "cohort", "dataset", "target_identity", "target_rank", "selection_origin",
        "base_sampling_seed", "eligibility_status", "train_start", "train_end",
        "val_start", "val_end", "test_start", "test_end", "n_train_days",
        "n_val_days", "n_test_days", "source_universe_id", "source_data_path",
        "source_data_sha256", "protocol_version", "selection_pool_size",
    ]

    def _excel_value(value):
        if isinstance(value, (list, tuple)):
            return ",".join("" if part is None else str(part) for part in value)
        if isinstance(value, dict):
            return json.dumps(value, sort_keys=True)
        return value

    def write_sheet(wb, name, rows, cols):
        ws = wb.create_sheet(name)
        fill = PatternFill("solid", fgColor="1F4E79")
        font = Font(bold=True, color="FFFFFF")
        ws.append(list(cols))
        for cell in ws[1]:
            cell.font = font
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in rows:
            ws.append([_excel_value(row.get(col)) for col in cols])
        ws.freeze_panes = "A2"
        last = get_column_letter(len(cols))
        ws.auto_filter.ref = f"A1:{last}{max(1, len(rows)+1)}"
        for i, col in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(col)) + 2, 12), 42)

    written = []
    method_rank = {m: i for i, m in enumerate(METHODS)}
    for dataset, meta in DATASETS.items():
        identity = _identity(dataset)
        calendar = calendar_for(dataset)
        raw = _iter_done(dataset)
        if not raw:
            raise OverlayStop("BLOCKED_SITA_PROTOCOL", f"{dataset} has no DONE cells")
        bad = [row for row in raw if int(row.get("n_test") or 0) != 180]
        if bad:
            raise OverlayStop("BLOCKED_SITA_PROTOCOL", f"{dataset} non-180 n_test={len(bad)}")
        selection = json.loads(_selection_path(dataset).read_text(encoding="utf-8")) if _selection_path(dataset).is_file() else {}
        selected = ",".join(item["identity"] for item in selection.get("selected", []))
        rows = []
        for item in raw:
            method = item.get("method")
            h = int(item["horizon"])
            seed = int(item["seed"])
            first = item.get("first_target_date") or calendar.test_start
            last = item.get("last_target_date") or calendar.test_end
            if first != calendar.test_start or last != calendar.test_end:
                raise OverlayStop("BLOCKED_SITA_PROTOCOL", f"{dataset} {method} h={h} dates {first}..{last}")
            store, sku = identity.split("/", 1)
            row = {key: None for key in CELL_COLUMNS}
            sources_raw = item.get("selected_sources")
            if isinstance(sources_raw, (list, tuple)):
                sources_raw = ",".join(str(part) for part in sources_raw)
            row.update(
                {
                    "cohort": COHORT,
                    "cohort_label": COHORT_LABEL,
                    "dataset": dataset,
                    "store_id": store,
                    "item_id": sku,
                    "target_identity": identity,
                    "target_rank": 1,
                    "selection_origin": SELECTION_ORIGIN,
                    "method": method,
                    "horizon": h,
                    "seed": seed,
                    "n_test": 180,
                    "rmse": item.get("rmse"),
                    "smape": item.get("smape"),
                    "status": "COMPLETED",
                    "first_target_date": first,
                    "last_target_date": last,
                    "train_start": calendar.train_start,
                    "train_end": calendar.train_end,
                    "val_start": calendar.val_start,
                    "val_end": calendar.val_end,
                    "n_train_days": TRAIN_DAYS,
                    "n_val_days": VAL_DAYS,
                    "n_test_days": TEST_DAYS,
                    "deterministic_reuse": method == "Seasonal Naive",
                    "prediction_source": "computed" if method == "Seasonal Naive" else "trained",
                    "prediction_finite_count": item.get("prediction_finite_count") or 180,
                    "target_finite_count": item.get("target_finite_count") or 180,
                    "selected_sources": sources_raw or (selected if method == "SITA" else None),
                    "cell_key": item.get("cell_key"),
                    "protocol_sha256": item.get("protocol_sha256"),
                    "source_universe_id": selection.get("source_universe_id"),
                    "source_data_path": meta["source_path"],
                    "direct_horizon": "y[forecast_origin + horizon]",
                    "trained_at": item.get("trained_at"),
                    "result_dir": item.get("_done"),
                    "canary_cell": False,
                    "cache_status": "PAPER_FORMAL_1_180",
                }
            )
            sources = str(row["selected_sources"] or "").split(",") if row["selected_sources"] else []
            for i, src in enumerate(sources[:3], 1):
                row[f"source_{i}"] = src
                if method == "SITA" and selection.get("selected"):
                    rec = selection["selected"][i - 1] if i - 1 < len(selection["selected"]) else {}
                    row[f"source_{i}_distance"] = rec.get("distance")
                    row[f"source_{i}_weight"] = rec.get("weight")
            if method == "MSML-TL":
                msml_sources = item.get("msml_selected_sources")
                if isinstance(msml_sources, (list, tuple)):
                    msml_sources = ",".join(str(part) for part in msml_sources)
                row["msml_selected_sources"] = msml_sources
                row["msml_formal_bl8"] = item.get("msml_formal_bl8")
                row["msml_used_sita_top3"] = item.get("msml_used_sita_top3")
                if not row.get("selected_sources") and msml_sources:
                    row["selected_sources"] = msml_sources
                    for i, src in enumerate(str(msml_sources).split(",")[:3], 1):
                        row[f"source_{i}"] = src
            rows.append(row)
        rows.sort(key=lambda r: (method_rank.get(r["method"], 99), r["horizon"], r["seed"]))
        if len(rows) != 135:
            raise OverlayStop("BLOCKED_SITA_PROTOCOL", f"{dataset} expected 135 cells, got {len(rows)}")
        from collections import defaultdict
        import statistics

        groups = defaultdict(list)
        for row in rows:
            groups[(row["method"], row["horizon"])].append(row)
        seed_rows = []
        for method in METHODS:
            for h in HORIZONS:
                items = groups[(method, h)]
                rmses = [float(x["rmse"]) for x in items]
                smapes = [float(x["smape"]) for x in items]
                seed_rows.append(
                    {
                        "cohort": COHORT,
                        "method": method,
                        "store_id": identity.split("/")[0],
                        "target_identity": identity,
                        "target_rank": 1,
                        "selection_origin": SELECTION_ORIGIN,
                        "horizon": h,
                        "n_seeds": len(items),
                        "rmse_mean": statistics.fmean(rmses),
                        "rmse_std": statistics.stdev(rmses) if len(rmses) > 1 else 0.0,
                        "smape_mean": statistics.fmean(smapes),
                        "smape_std": statistics.stdev(smapes) if len(smapes) > 1 else 0.0,
                        "n_test": 180,
                        "first_target_date": calendar.test_start,
                        "last_target_date": calendar.test_end,
                        "selected_sources": items[0].get("selected_sources"),
                        "deterministic_reuse": method == "Seasonal Naive",
                    }
                )
        coverage = []
        for method in METHODS:
            done = sum(1 for row in rows if row["method"] == method)
            coverage.append(
                {
                    "cohort": COHORT,
                    "method": method,
                    "n_targets_expected": 1,
                    "n_cells_expected": 15,
                    "n_cells_completed": done,
                    "n_cells_other": 0,
                    "complete": done == 15,
                }
            )
        targets = [
            {
                "cohort": COHORT,
                "dataset": dataset,
                "target_identity": identity,
                "target_rank": 1,
                "selection_origin": SELECTION_ORIGIN,
                "base_sampling_seed": None,
                "eligibility_status": "paper_formal_target_180",
                "train_start": calendar.train_start,
                "train_end": calendar.train_end,
                "val_start": calendar.val_start,
                "val_end": calendar.val_end,
                "test_start": calendar.test_start,
                "test_end": calendar.test_end,
                "n_train_days": TRAIN_DAYS,
                "n_val_days": VAL_DAYS,
                "n_test_days": TEST_DAYS,
                "source_universe_id": selection.get("source_universe_id"),
                "source_data_path": meta["source_path"],
                "source_data_sha256": None,
                "protocol_version": "paper_formal_1_180",
                "selection_pool_size": selection.get("legal_count"),
            }
        ]
        readme = [
            {"field": "dataset", "value": dataset},
            {"field": "experiment", "value": f"{dataset} paper formal 1-target restated on 180-day robustness protocol"},
            {"field": "target_identity", "value": identity},
            {"field": "cohort", "value": COHORT},
            {"field": "n_test", "value": 180},
            {"field": "test_start", "value": calendar.test_start},
            {"field": "test_end", "value": calendar.test_end},
            {"field": "n_cells_total", "value": len(rows)},
            {"field": "old_paper1target_kept", "value": str((XLSX_DIR / meta["old_xlsx"]).relative_to(REPO_ROOT))},
            {"field": "not_merged_from_170", "value": True},
            {"field": "written_at", "value": _utc()},
        ]
        wb = Workbook()
        wb.remove(wb.active)
        write_sheet(wb, "cell_raw", rows, CELL_COLUMNS)
        write_sheet(wb, "seed_mean", seed_rows, seed_cols)
        write_sheet(wb, "coverage", coverage, cov_cols)
        write_sheet(wb, "targets", targets, tgt_cols)
        write_sheet(wb, "readme", readme, ["field", "value"])
        out = XLSX_DIR / meta["xlsx"]
        wb.save(out)
        written.append(out)
        print(out.relative_to(REPO_ROOT), "rows", len(rows))
    return written


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--phase",
        choices=("audit", "prepare", "stage-a", "sita", "baselines", "workbook"),
        required=True,
    )
    parser.add_argument("--dataset", choices=("D4", "D5", "D6", "all"), default="all")
    parser.add_argument("--methods", default="", help="baseline subset, comma-separated")
    args = parser.parse_args()
    datasets = list(DATASETS) if args.dataset == "all" else [args.dataset]
    methods = tuple(item.strip() for item in args.methods.split(",") if item.strip()) or None
    try:
        if args.phase == "audit":
            payload = phase_audit()
            print(json.dumps({"status": payload["status"], "path": str((PLAN_DIR / "00_audit.md").relative_to(REPO_ROOT))}, indent=2))
            return 0
        if args.phase == "workbook":
            phase_workbook()
            return 0
        for dataset in datasets:
            if args.phase == "prepare":
                print(json.dumps(phase_prepare(dataset), indent=2, sort_keys=True))
            elif args.phase == "stage-a":
                print(json.dumps(phase_stage_a(dataset), indent=2, sort_keys=True))
            elif args.phase == "sita":
                print(json.dumps(phase_sita(dataset), indent=2, sort_keys=True))
            elif args.phase == "baselines":
                print(json.dumps(phase_baselines(dataset, methods=methods), indent=2, sort_keys=True))
        return 0
    except OverlayStop as exc:
        print(f"{exc.token}: {exc.detail or exc}", file=sys.stderr)
        return 2
    except ProtocolError as exc:
        print(f"BLOCKED_SITA_PROTOCOL: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
