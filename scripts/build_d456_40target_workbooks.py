#!/usr/bin/env python3
"""Build D4/D5/D6 40-target workbooks in the 70-target BL+SITA sheet format.

Keeps the swapped fixed-40 set: 39 original instances plus one paper formal
target restated on the 180-day protocol.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RESULT = ROOT / "实验结果"
MANIFEST = RESULT / "sita_d456_fixed40_headline_paper180swap" / "fixed40_manifest.csv"
METHOD_ORDER = (
    "SITA",
    "Seasonal Naive",
    "LightGBM",
    "LSTM",
    "DLinear",
    "TCN",
    "PatchTST",
    "TimesNet",
    "MSML-TL",
)
SWAPS = {
    "D4": {"drop": "166/229", "add": "166/258", "paper": RESULT / "D4_paper1target_BL_SITA_180.xlsx"},
    "D5": {"drop": "48/414750", "add": "48/1159414", "paper": RESULT / "D5_paper1target_BL_SITA_180.xlsx"},
    "D6": {"drop": "CA_1/FOODS_3_255", "add": "CA_1/FOODS_3_080", "paper": RESULT / "D6_paper1target_BL_SITA_180.xlsx"},
}


def utc_now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def write_sheet(wb: Workbook, name: str, frame: pd.DataFrame) -> None:
    ws = wb.create_sheet(name)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    columns = list(frame.columns)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for rec in frame.itertuples(index=False, name=None):
        row = []
        for value in rec:
            if value is None or (isinstance(value, float) and pd.isna(value)):
                row.append(None)
            elif isinstance(value, (np.integer,)):
                row.append(int(value))
            elif isinstance(value, (np.floating,)):
                row.append(float(value))
            elif isinstance(value, np.bool_):
                row.append(bool(value))
            else:
                row.append(value)
        ws.append(row)
    n_rows = max(1, len(frame) + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{n_rows}"
    ws.freeze_panes = "A2"
    for index, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = min(max(len(column) + 2, 12), 42)


def method_horizon_summary(seed: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (cohort, method, horizon), grp in seed.groupby(["cohort", "method", "horizon"], dropna=False):
        rows.append(_summary_row(str(cohort), method, horizon, grp))
    for (method, horizon), grp in seed.groupby(["method", "horizon"], dropna=False):
        rows.append(_summary_row("ALL_40", method, horizon, grp))
    out = pd.DataFrame(rows)
    out["method_ord"] = out["method"].map(lambda name: METHOD_ORDER.index(name) if name in METHOD_ORDER else 99)
    cohort_ord = {"OLD_ADDITIONAL_10": 0, "NEW_RANDOM_60": 1, "PAPER_FORMAL_1_180": 2, "ALL_40": 3}
    out["cohort_ord"] = out["cohort"].map(lambda name: cohort_ord.get(name, 9))
    out = out.sort_values(["cohort_ord", "method_ord", "horizon"]).drop(columns=["method_ord", "cohort_ord"])
    return out.reset_index(drop=True)


def _summary_row(cohort: str, method, horizon, grp: pd.DataFrame) -> dict:
    rmse = pd.to_numeric(grp["rmse_mean"], errors="coerce")
    smape = pd.to_numeric(grp["smape_mean"], errors="coerce")
    return {
        "cohort": cohort,
        "method": method,
        "horizon": int(horizon),
        "n_targets": int(len(grp)),
        "rmse_mean": float(rmse.mean()),
        "rmse_median": float(rmse.median()),
        "rmse_std": float(rmse.std(ddof=0)) if len(grp) > 1 else 0.0,
        "smape_mean": float(smape.mean()),
        "smape_median": float(smape.median()),
        "smape_std": float(smape.std(ddof=0)) if len(grp) > 1 else 0.0,
    }


def coverage_rows(cell: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for cohort, grp in cell.groupby("cohort", dropna=False):
        n_targets = int(grp["target_identity"].nunique())
        for method in METHOD_ORDER:
            sub = grp.loc[grp["method"] == method]
            expected = n_targets * 3 * 5
            completed = int((sub["status"].astype(str).isin(["COMPLETED", "SKIP"])).sum()) if not sub.empty else 0
            other = int((~sub["status"].astype(str).isin(["COMPLETED", "SKIP"])).sum()) if not sub.empty else 0
            rows.append(
                {
                    "cohort": cohort,
                    "method": method,
                    "n_targets_expected": n_targets,
                    "n_cells_expected": expected,
                    "n_cells_completed": completed,
                    "n_cells_other": other,
                    "complete": completed == expected and other == 0 and expected > 0,
                }
            )
    n_all = int(cell["target_identity"].nunique())
    for method in METHOD_ORDER:
        sub = cell.loc[cell["method"] == method]
        expected = n_all * 3 * 5
        completed = int((sub["status"].astype(str).isin(["COMPLETED", "SKIP"])).sum())
        other = int((~sub["status"].astype(str).isin(["COMPLETED", "SKIP"])).sum())
        rows.append(
            {
                "cohort": "ALL_40",
                "method": method,
                "n_targets_expected": n_all,
                "n_cells_expected": expected,
                "n_cells_completed": completed,
                "n_cells_other": other,
                "complete": completed == expected and other == 0,
            }
        )
    return pd.DataFrame(rows)


def build_one(dataset: str, manifest: pd.DataFrame) -> Path:
    src70 = RESULT / f"{dataset}_70target_BL_SITA.xlsx"
    paper_path = SWAPS[dataset]["paper"]
    keep = manifest.loc[manifest["dataset"] == dataset].copy()
    if len(keep) != 40:
        raise RuntimeError(f"{dataset}: manifest has {len(keep)} rows, expected 40")
    orig = keep.loc[keep["cohort"] != "PAPER_FORMAL_1_180"]
    paper_man = keep.loc[keep["cohort"] == "PAPER_FORMAL_1_180"]
    if len(orig) != 39 or len(paper_man) != 1:
        raise RuntimeError(f"{dataset}: expected 39 original + 1 paper, got {len(orig)}+{len(paper_man)}")

    cell70 = pd.read_excel(src70, "cell_raw")
    seed70 = pd.read_excel(src70, "seed_mean")
    tgt70 = pd.read_excel(src70, "targets")
    paper_xl = pd.ExcelFile(paper_path)
    cell_p = pd.read_excel(paper_path, "cell_raw")
    seed_p = pd.read_excel(paper_path, "seed_mean")
    tgt_p = pd.read_excel(paper_path, "targets") if "targets" in paper_xl.sheet_names else pd.DataFrame()

    orig_keys = set(zip(orig["cohort"].astype(str), orig["target_identity"].astype(str)))
    rank_map = {
        (str(r.cohort), str(r.target_identity)): int(r.list_rank)
        for r in keep.itertuples(index=False)
    }

    cell_o = cell70.loc[
        [(str(c), str(i)) in orig_keys for c, i in zip(cell70["cohort"], cell70["target_identity"])]
    ].copy()
    seed_o = seed70.loc[
        [(str(c), str(i)) in orig_keys for c, i in zip(seed70["cohort"], seed70["target_identity"])]
    ].copy()
    tgt_o = tgt70.loc[
        [(str(c), str(i)) in orig_keys for c, i in zip(tgt70["cohort"], tgt70["target_identity"])]
    ].copy()

    def apply_rank(frame: pd.DataFrame) -> pd.DataFrame:
        out = frame.copy()
        out["target_rank"] = [
            rank_map[(str(c), str(i))] for c, i in zip(out["cohort"], out["target_identity"])
        ]
        return out

    cell = pd.concat([apply_rank(cell_o), apply_rank(cell_p)], ignore_index=True)
    seed = pd.concat([apply_rank(seed_o), apply_rank(seed_p)], ignore_index=True)
    tgt = pd.concat([apply_rank(tgt_o), apply_rank(tgt_p)], ignore_index=True)
    tgt = tgt.sort_values("target_rank").reset_index(drop=True)

    # stable sort like original builder: identity then method/horizon/seed
    cell = cell.sort_values(["target_rank", "method", "horizon", "seed"], kind="mergesort").reset_index(drop=True)
    seed = seed.sort_values(["target_rank", "method", "horizon"], kind="mergesort").reset_index(drop=True)

    n_targets = int(tgt["target_identity"].nunique())
    if n_targets != 40:
        raise RuntimeError(f"{dataset}: unique identities={n_targets}")
    expected_cells = 40 * 9 * 3 * 5
    if len(cell) != expected_cells:
        raise RuntimeError(f"{dataset}: cell_raw rows={len(cell)} expected {expected_cells}")
    if len(seed) != 40 * 9 * 3:
        raise RuntimeError(f"{dataset}: seed_mean rows={len(seed)} expected {1080}")
    if not (pd.to_numeric(cell["n_test"], errors="coerce") == 180).all():
        raise RuntimeError(f"{dataset}: n_test is not all 180")

    summary = method_horizon_summary(seed)
    coverage = coverage_rows(cell)
    if not coverage.loc[coverage["cohort"] == "ALL_40", "complete"].all():
        raise RuntimeError(f"{dataset}: ALL_40 coverage incomplete")

    n_old = int((tgt["cohort"] == "OLD_ADDITIONAL_10").sum())
    n_new = int((tgt["cohort"] == "NEW_RANDOM_60").sum())
    n_paper = int((tgt["cohort"] == "PAPER_FORMAL_1_180").sum())
    readme = pd.DataFrame(
        [
            {"field": "dataset", "value": dataset},
            {"field": "experiment", "value": f"{dataset} 40-target BL+SITA after one paper-target swap"},
            {"field": "format", "value": "same sheets as D*_70target_BL_SITA.xlsx"},
            {"field": "n_targets_total", "value": 40},
            {"field": "n_targets_from_70target_workbook", "value": 39},
            {"field": "n_targets_old_additional_10", "value": n_old},
            {"field": "n_targets_new_random_60", "value": n_new},
            {"field": "n_targets_paper_180", "value": n_paper},
            {"field": "replaced_out", "value": SWAPS[dataset]["drop"]},
            {"field": "replaced_in", "value": SWAPS[dataset]["add"]},
            {"field": "paper_source", "value": str(paper_path.relative_to(ROOT))},
            {"field": "source_70target", "value": str(src70.relative_to(ROOT))},
            {"field": "horizons", "value": "1, 3, 5"},
            {"field": "seeds", "value": "42, 43, 44, 45, 46"},
            {"field": "n_test", "value": 180},
            {"field": "evaluation_index", "value": "direct-h; every test day is a label; n_test=180"},
            {"field": "methods", "value": ", ".join(METHOD_ORDER)},
            {"field": "n_cells_total", "value": len(cell)},
            {"field": "n_seed_mean_rows", "value": len(seed)},
            {"field": "manifest", "value": str(MANIFEST.relative_to(ROOT))},
            {"field": "written_at", "value": utc_now()},
        ]
    )

    out_path = RESULT / f"{dataset}_40target_BL_SITA.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "cell_raw", cell)
    write_sheet(wb, "seed_mean", seed)
    write_sheet(wb, "method_horizon_summary", summary)
    write_sheet(wb, "coverage", coverage)
    write_sheet(wb, "targets", tgt)
    write_sheet(wb, "readme", readme)
    wb.save(out_path)
    return out_path


def main() -> int:
    manifest = pd.read_csv(MANIFEST)
    for dataset in ("D4", "D5", "D6"):
        path = build_one(dataset, manifest)
        print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
