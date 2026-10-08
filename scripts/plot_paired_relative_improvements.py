#!/usr/bin/env python3
"""Create paired relative-improvement figures for the frozen SITA A-E ablation.

The script reads condition-level five-seed means from the frozen workbook, derives
four paired contrasts, validates the 12-condition contract, and writes separate
RMSE and sMAPE figures plus machine-readable data and captions.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import math
import os
import re
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from statistics import mean, median

# Keep Matplotlib's font/cache writes outside the repository.  This also makes
# the documented command work without a shell-specific environment preamble.
os.environ.setdefault("MPLCONFIGDIR", os.path.join(tempfile.gettempdir(), "sita_paired_mpl_cache"))
os.environ.setdefault("XDG_CACHE_HOME", os.path.join(tempfile.gettempdir(), "sita_paired_xdg_cache"))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.lines import Line2D
from openpyxl import load_workbook


EXPECTED_DATASETS = ["D1", "D4", "D5", "D6"]
EXPECTED_HORIZONS = [1, 3, 5]
EXPECTED_ARMS = ["A", "B", "C", "D", "E"]
CONTRASTS = [
    {
        "key": "C_vs_A",
        "improved_arm": "C",
        "baseline_arm": "A",
        "label": "C vs A\nFused init.",
        "question": "Fused source-informed initialization vs native initialization",
    },
    {
        "key": "C_vs_B",
        "improved_arm": "C",
        "baseline_arm": "B",
        "label": "C vs B\nTrainability",
        "question": "Trainable vs frozen convolutional backbone",
    },
    {
        "key": "D_vs_C",
        "improved_arm": "D",
        "baseline_arm": "C",
        "label": "D vs C\nRevIN",
        "question": "RevIN contribution",
    },
    {
        "key": "C_vs_E",
        "improved_arm": "C",
        "baseline_arm": "E",
        "label": "C vs E\nMulti-source",
        "question": "Multi-source fused prior vs best-single-source prior",
    },
]
METRICS = {
    "RMSE": {"summary_column": "RMSE_Mean", "raw_column": "rmse_original"},
    "sMAPE": {"summary_column": "sMAPE_Mean", "raw_column": "smape_original"},
}
DATASET_COLORS = {
    "D1": "#E69F00",  # Okabe-Ito orange
    "D4": "#56B4E9",  # Okabe-Ito sky blue
    "D5": "#009E73",  # Okabe-Ito green
    "D6": "#0072B2",  # Okabe-Ito blue
}
HORIZON_MARKERS = {1: "o", 3: "s", 5: "^"}
CONDITION_ORDER = [(d, h) for d in EXPECTED_DATASETS for h in EXPECTED_HORIZONS]
TIE_TOLERANCE = 1e-12


class FreezeValidationError(RuntimeError):
    """Raised when the frozen input cannot satisfy the plotting contract."""


def normalized(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def value_rows(workbook, sheet_name: str) -> list[dict[str, object]]:
    if sheet_name not in workbook.sheetnames:
        return []
    rows = list(workbook[sheet_name].values)
    if not rows:
        return []
    headers = list(rows[0])
    records = []
    for row in rows[1:]:
        if any(value is not None for value in row):
            records.append(dict(zip(headers, row)))
    return records


def detect_column(headers: list[object], candidates: list[str]) -> object | None:
    by_normalized = {normalized(header): header for header in headers if header is not None}
    for candidate in candidates:
        if normalized(candidate) in by_normalized:
            return by_normalized[normalized(candidate)]
    return None


def require_columns(records: list[dict[str, object]], column_map: dict[str, list[str]], source: str):
    headers = list(records[0]) if records else []
    detected = {}
    missing = []
    for logical_name, candidates in column_map.items():
        detected[logical_name] = detect_column(headers, candidates)
        if detected[logical_name] is None:
            missing.append(f"{logical_name}={candidates}")
    if missing:
        raise FreezeValidationError(f"Missing columns in {source}: {', '.join(missing)}")
    return detected


def as_number(value: object, field: str, row_number: int) -> float:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise FreezeValidationError(f"Missing numeric value {field} at row {row_number}")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise FreezeValidationError(f"Non-numeric value {value!r} in {field} at row {row_number}") from exc
    if not math.isfinite(result):
        raise FreezeValidationError(f"Non-finite value in {field} at row {row_number}")
    return result


def load_condition_means(workbook) -> tuple[dict[tuple[str, int, str], dict[str, object]], str]:
    """Load existing five-seed condition means, falling back to raw aggregation."""

    summary_records = value_rows(workbook, "11_Ablation_Summary")
    if summary_records:
        columns = require_columns(
            summary_records,
            {
                "dataset": ["Dataset", "dataset"],
                "horizon": ["Horizon", "horizon"],
                "arm": ["Arm", "arm"],
                "n_seeds": ["N_Seeds", "n_seeds", "Seeds"],
                "rmse": ["RMSE_Mean", "rmse_mean", "rmse"],
                "smape": ["sMAPE_Mean", "smape_mean", "smape"],
            },
            "11_Ablation_Summary",
        )
        means = {}
        for row_number, row in enumerate(summary_records, start=2):
            dataset = str(row[columns["dataset"]]).strip().upper()
            horizon = int(float(row[columns["horizon"]]))
            arm = str(row[columns["arm"]]).strip().upper()
            n_seeds = int(float(row[columns["n_seeds"]]))
            if n_seeds != 5:
                raise FreezeValidationError(
                    f"Expected five-seed condition means, found N_Seeds={n_seeds} at row {row_number}"
                )
            key = (dataset, horizon, arm)
            if key in means:
                raise FreezeValidationError(f"Duplicate condition-arm in summary: {key}")
            means[key] = {
                "dataset": dataset,
                "horizon": horizon,
                "arm": arm,
                "RMSE": as_number(row[columns["rmse"]], "RMSE_Mean", row_number),
                "sMAPE": as_number(row[columns["smape"]], "sMAPE_Mean", row_number),
                "seed_aggregation": "five-seed condition-level mean from 11_Ablation_Summary",
                "source_sheet": "11_Ablation_Summary",
            }
        return validate_condition_means(means), "11_Ablation_Summary"

    raw_records = value_rows(workbook, "05_Ablation_Raw")
    if not raw_records:
        raise FreezeValidationError("Neither 11_Ablation_Summary nor 05_Ablation_Raw is available")
    columns = require_columns(
        raw_records,
        {
            "dataset": ["dataset", "Dataset"],
            "horizon": ["horizon", "Horizon"],
            "seed": ["seed", "Seed"],
            "arm": ["arm", "Arm"],
            "rmse": ["rmse_original", "RMSE_Mean", "rmse"],
            "smape": ["smape_original", "sMAPE_Mean", "smape"],
            "status": ["status", "Status"],
            "metric_space": ["metric_space", "Metric_Space"],
        },
        "05_Ablation_Raw",
    )
    grouped: defaultdict[tuple[str, int, str], dict[str, list[float]]] = defaultdict(
        lambda: {"RMSE": [], "sMAPE": [], "seeds": []}
    )
    for row_number, row in enumerate(raw_records, start=2):
        status = str(row[columns["status"]]).strip().upper()
        metric_space = str(row[columns["metric_space"]]).strip()
        if status != "OK":
            raise FreezeValidationError(f"Non-OK raw row at row {row_number}: {status!r}")
        if metric_space != "original_sales_space":
            raise FreezeValidationError(f"Unexpected metric space at row {row_number}: {metric_space!r}")
        dataset = str(row[columns["dataset"]]).strip().upper()
        horizon = int(float(row[columns["horizon"]]))
        seed = int(float(row[columns["seed"]]))
        arm = str(row[columns["arm"]]).strip().upper()
        key = (dataset, horizon, arm)
        grouped[key]["RMSE"].append(as_number(row[columns["rmse"]], "rmse_original", row_number))
        grouped[key]["sMAPE"].append(as_number(row[columns["smape"]], "smape_original", row_number))
        grouped[key]["seeds"].append(seed)
    means = {}
    for key, values in grouped.items():
        if len(values["seeds"]) != 5 or len(set(values["seeds"])) != 5:
            raise FreezeValidationError(f"Expected five unique seeds for {key}, found {values['seeds']}")
        dataset, horizon, arm = key
        means[key] = {
            "dataset": dataset,
            "horizon": horizon,
            "arm": arm,
            "RMSE": mean(values["RMSE"]),
            "sMAPE": mean(values["sMAPE"]),
            "seed_aggregation": "mean of five raw seed rows from 05_Ablation_Raw",
            "source_sheet": "05_Ablation_Raw",
        }
    return validate_condition_means(means), "05_Ablation_Raw"


def validate_condition_means(means: dict[tuple[str, int, str], dict[str, object]]):
    expected_keys = {
        (dataset, horizon, arm)
        for dataset in EXPECTED_DATASETS
        for horizon in EXPECTED_HORIZONS
        for arm in EXPECTED_ARMS
    }
    # ``means`` is already keyed by the condition-arm tuple; duplicate summary
    # rows are rejected while loading, and duplicate raw seeds are checked
    # before aggregation.  Do not pass this mapping to Counter, which would
    # interpret each row dictionary as a non-unit count.
    duplicate_keys = []
    actual_keys = set(means)
    missing = sorted(expected_keys - actual_keys)
    unexpected = sorted(actual_keys - expected_keys)
    if duplicate_keys or missing or unexpected:
        raise FreezeValidationError(
            f"Condition validation failed: duplicates={duplicate_keys}, missing={missing}, unexpected={unexpected}"
        )
    for key, row in means.items():
        for metric in METRICS:
            if not math.isfinite(float(row[metric])):
                raise FreezeValidationError(f"Non-finite {metric} for {key}")
    return means


def audit_raw_sheet(workbook) -> dict[str, object]:
    """Audit raw row completeness independently of the selected summary sheet."""

    records = value_rows(workbook, "05_Ablation_Raw")
    if not records:
        return {"rows": None, "duplicates": [], "missing_values": None, "status_counts": {}, "metric_space_counts": {}}
    columns = require_columns(
        records,
        {
            "dataset": ["dataset", "Dataset"],
            "horizon": ["horizon", "Horizon"],
            "seed": ["seed", "Seed"],
            "arm": ["arm", "Arm"],
            "rmse": ["rmse_original", "RMSE_Mean", "rmse"],
            "smape": ["smape_original", "sMAPE_Mean", "smape"],
            "status": ["status", "Status"],
            "metric_space": ["metric_space", "Metric_Space"],
        },
        "05_Ablation_Raw",
    )
    keys = []
    missing_values = 0
    status_counts = Counter()
    metric_space_counts = Counter()
    for row in records:
        key = (
            str(row[columns["dataset"]]).strip().upper(),
            int(float(row[columns["horizon"]])),
            int(float(row[columns["seed"]])),
            str(row[columns["arm"]]).strip().upper(),
        )
        keys.append(key)
        status_counts[str(row[columns["status"]]).strip()] += 1
        metric_space_counts[str(row[columns["metric_space"]]).strip()] += 1
        if row[columns["rmse"]] is None or row[columns["smape"]] is None:
            missing_values += 1
    duplicate_counts = Counter(keys)
    duplicates = sorted((key, count) for key, count in duplicate_counts.items() if count != 1)
    return {
        "rows": len(records),
        "duplicates": duplicates,
        "missing_values": missing_values,
        "status_counts": dict(status_counts),
        "metric_space_counts": dict(metric_space_counts),
    }


def load_statistics(workbook) -> dict[tuple[str, str], dict[str, object]]:
    """Read existing frozen contrast CIs and tests without recomputing them."""

    statistics: dict[tuple[str, str], dict[str, object]] = defaultdict(dict)
    evidence_rows = value_rows(workbook, "16_Ablation_Contrast_Evidence")
    if evidence_rows:
        columns = require_columns(
            evidence_rows,
            {
                "contrast": ["Contrast"],
                "grade": ["Evidence_Grade"],
                "claim": ["Paper_Safe_Claim"],
                "caveat": ["Caveat"],
            },
            "16_Ablation_Contrast_Evidence",
        )
        for row in evidence_rows:
            contrast = str(row[columns["contrast"]]).strip()
            for metric in METRICS:
                statistics[(contrast, metric)].update(
                    {
                        "evidence_grade": row[columns["grade"]],
                        "paper_safe_claim": row[columns["claim"]],
                        "caveat": row[columns["caveat"]],
                    }
                )

    test_rows = value_rows(workbook, "17_Ablation_Statistical_Tests")
    if test_rows:
        columns = require_columns(
            test_rows,
            {
                "contrast": ["Contrast"],
                "metric": ["Metric"],
                "n_conditions": ["N_Conditions"],
                "mean": ["Mean_Relative_Improvement_pct"],
                "median": ["Median_Relative_Improvement_pct"],
                "ci_low": ["Bootstrap_95CI_Low_pct"],
                "ci_high": ["Bootstrap_95CI_High_pct"],
                "wilcoxon_w": ["Wilcoxon_W"],
                "wilcoxon_p": ["Wilcoxon_Raw_p"],
                "holm_p": ["Holm_Adjusted_p"],
                "interpretation": ["Interpretation"],
            },
            "17_Ablation_Statistical_Tests",
        )
        for row_number, row in enumerate(test_rows, start=2):
            contrast = str(row[columns["contrast"]]).strip()
            metric_value = str(row[columns["metric"]]).strip().upper()
            metric = "sMAPE" if metric_value == "SMAPE" else metric_value
            if metric not in METRICS:
                continue
            statistics[(contrast, metric)].update(
                {
                    "n_conditions": int(float(row[columns["n_conditions"]])),
                    "frozen_mean": as_number(row[columns["mean"]], "Mean_Relative_Improvement_pct", row_number),
                    "frozen_median": as_number(row[columns["median"]], "Median_Relative_Improvement_pct", row_number),
                    "ci_low": as_number(row[columns["ci_low"]], "Bootstrap_95CI_Low_pct", row_number),
                    "ci_high": as_number(row[columns["ci_high"]], "Bootstrap_95CI_High_pct", row_number),
                    "wilcoxon_w": row[columns["wilcoxon_w"]],
                    "wilcoxon_raw_p": as_number(row[columns["wilcoxon_p"]], "Wilcoxon_Raw_p", row_number),
                    "holm_adjusted_p": as_number(row[columns["holm_p"]], "Holm_Adjusted_p", row_number),
                    "interpretation": row[columns["interpretation"]],
                    "bootstrap_source": "17_Ablation_Statistical_Tests",
                }
            )
    return statistics


def derive_relative_improvements(means: dict[tuple[str, int, str], dict[str, object]]):
    records = {metric: [] for metric in METRICS}
    for contrast in CONTRASTS:
        for dataset, horizon in CONDITION_ORDER:
            improved = means[(dataset, horizon, contrast["improved_arm"])]
            baseline = means[(dataset, horizon, contrast["baseline_arm"])]
            for metric in METRICS:
                denominator = float(baseline[metric])
                if denominator == 0:
                    raise FreezeValidationError(
                        f"Zero baseline denominator for {contrast['key']} {metric} {dataset}-H{horizon}"
                    )
                ri = (denominator - float(improved[metric])) / denominator * 100.0
                records[metric].append(
                    {
                        "contrast": contrast["key"],
                        "contrast_label": contrast["label"].replace("\n", " / "),
                        "research_question": contrast["question"],
                        "dataset": dataset,
                        "horizon": horizon,
                        "condition": f"{dataset}-H{horizon}",
                        "improved_arm": contrast["improved_arm"],
                        "baseline_arm": contrast["baseline_arm"],
                        "baseline_value": denominator,
                        "improved_value": float(improved[metric]),
                        "relative_improvement_pct": ri,
                        "metric": metric,
                        "metric_space": "original_sales_space" if metric == "RMSE" else "percent",
                        "seed_aggregation": means[(dataset, horizon, contrast["improved_arm"])] ["seed_aggregation"],
                    }
                )
    return records


def summarize(records: list[dict[str, object]], statistics: dict[tuple[str, str], dict[str, object]], metric: str):
    summaries = []
    for contrast in CONTRASTS:
        subset = [row for row in records if row["contrast"] == contrast["key"]]
        values = [float(row["relative_improvement_pct"]) for row in subset]
        wins = sum(value > TIE_TOLERANCE for value in values)
        ties = sum(abs(value) <= TIE_TOLERANCE for value in values)
        losses = sum(value < -TIE_TOLERANCE for value in values)
        stat = statistics.get((contrast["key"], metric), {})
        summary = {
            "contrast": contrast["key"],
            "metric": metric,
            "n_conditions": len(values),
            "mean_RI_pct": mean(values),
            "median_RI_pct": median(values),
            "min_RI_pct": min(values),
            "max_RI_pct": max(values),
            "wins": wins,
            "ties": ties,
            "losses": losses,
            "bootstrap_95CI_low_pct": stat.get("ci_low"),
            "bootstrap_95CI_high_pct": stat.get("ci_high"),
            "wilcoxon_W": stat.get("wilcoxon_w"),
            "wilcoxon_raw_p": stat.get("wilcoxon_raw_p"),
            "holm_adjusted_p": stat.get("holm_adjusted_p"),
            "evidence_grade": stat.get("evidence_grade"),
            "paper_safe_claim": stat.get("paper_safe_claim"),
            "caveat": stat.get("caveat"),
        }
        if stat.get("frozen_mean") is not None:
            summary["derived_minus_frozen_mean_pct"] = summary["mean_RI_pct"] - stat["frozen_mean"]
        if len(values) != 12:
            raise FreezeValidationError(
                f"Expected 12 points for {metric}/{contrast['key']}, found {len(values)}"
            )
        summaries.append(summary)
    return summaries


def format_number(value: object, digits: int = 4) -> str:
    if value is None or value == "":
        return "NA"
    if isinstance(value, (float, int, np.floating, np.integer)):
        return f"{float(value):.{digits}f}"
    return str(value)


def write_csv(path: Path, rows: list[dict[str, object]]):
    if not rows:
        return
    fields = list(rows[0])
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_summary(path: Path, input_path: Path, sha256: str, source_sheet: str, raw_audit: dict[str, object], summaries):
    lines = [
        "# Paired Relative Improvements of Major Mechanism Contrasts",
        "",
        "## Provenance and validation",
        "",
        f"- Input workbook: `{input_path}`",
        f"- SHA256: `{sha256}`",
        f"- Selected data sheet: `{source_sheet}`",
        "- A–E arms: A, B, C, D, E; condition-level five-seed means",
        "- Conditions: D1/D4/D5/D6 × H1/H3/H5 = 12 paired conditions",
        "- RMSE metric space: `original_sales_space`; sMAPE unit: percent",
        "- Relative improvement: `(baseline arm - improved arm) / baseline arm × 100`",
        "- Existing bootstrap CIs and Wilcoxon/Holm values are copied from the frozen statistical sheet; no new statistical test was run.",
        f"- Raw audit: {raw_audit.get('rows')} rows; statuses={raw_audit.get('status_counts')}; metric_spaces={raw_audit.get('metric_space_counts')}; missing_values={raw_audit.get('missing_values')}; duplicate_seed_arm_keys={raw_audit.get('duplicates')}",
        "",
        "## Summary statistics",
        "",
    ]
    for metric in ["RMSE", "sMAPE"]:
        lines.extend(
            [
                f"### {metric}",
                "",
                "| Contrast | Mean RI (%) | Median RI (%) | Min | Max | Wins/Ties/Losses | Bootstrap 95% CI | Holm-adjusted p | Evidence grade |",
                "|---|---:|---:|---:|---:|---:|---:|---:|---|",
            ]
        )
        for row in [item for item in summaries if item["metric"] == metric]:
            ci = f"[{format_number(row['bootstrap_95CI_low_pct'], 2)}, {format_number(row['bootstrap_95CI_high_pct'], 2)}]"
            wtl = f"{row['wins']}/{row['ties']}/{row['losses']}"
            lines.append(
                f"| {row['contrast']} | {format_number(row['mean_RI_pct'], 4)} | {format_number(row['median_RI_pct'], 4)} | {format_number(row['min_RI_pct'], 4)} | {format_number(row['max_RI_pct'], 4)} | {wtl} | {ci} | {format_number(row['holm_adjusted_p'], 6)} | {row['evidence_grade'] or 'NA'} |"
            )
        lines.append("")
    lines.extend(
        [
            "## Frozen claim boundary",
            "",
            "- C vs A: fused source-informed initialization versus native initialization.",
            "- C vs B: trainable versus frozen convolutional backbone; do not relabel this as adaptation on/off.",
            "- D vs C: RevIN contribution; do not relabel this as domain alignment or source-target alignment.",
            "- C vs E: multi-source fused prior versus best-single-source prior; do not claim proven optimal source fusion.",
            "",
            "## Captions",
            "",
            "### English",
            "",
            "**Figure X. Paired relative improvements of major mechanism contrasts.** Each point denotes the condition-level relative improvement for one dataset–horizon pair (D1/D4/D5/D6 × H1/H3/H5). Improvements are computed pairwise between the compared ablation arms, so positive values indicate that the highlighted mechanism improves performance. C vs A isolates the contribution of fused source-informed initialization, C vs B tests whether the fused prior should remain trainable rather than frozen, D vs C measures the contribution of RevIN, and C vs E compares multi-source fused initialization against best-single-source initialization. Large diamonds denote mean relative improvement across the 12 conditions; vertical whiskers show the frozen bootstrap 95% confidence interval.",
            "",
            "### 中文",
            "",
            "**图 X. 主要机制对比的配对相对改进。** 每个点表示一个 dataset–horizon 配对条件（D1/D4/D5/D6 × H1/H3/H5）的条件级相对改进。改进值均在对应消融臂之间成对计算，因此大于 0 表示所标示机制带来性能改善。C vs A 分离融合源信息初始化相对于原生初始化的贡献；C vs B 检验融合先验下可训练卷积骨干相对于冻结骨干的影响；D vs C 衡量 RevIN 的贡献；C vs E 比较多源融合初始化与最佳单源初始化。大菱形表示 12 个条件上的平均相对改进，纵向误差线表示冻结统计表中的 bootstrap 95% 置信区间。",
        ]
    )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_captions(path: Path):
    path.write_text(
        """# Captions\n\n## English\n\n**Figure X. Paired relative improvements of major mechanism contrasts (RMSE / sMAPE).** Each point denotes the condition-level relative improvement for one dataset–horizon pair (D1/D4/D5/D6 × H1/H3/H5). Improvements are computed pairwise between the compared ablation arms, so positive values indicate that the highlighted mechanism improves performance. C vs A isolates the contribution of fused source-informed initialization, C vs B tests whether the fused prior should remain trainable rather than frozen, D vs C measures the contribution of RevIN, and C vs E compares multi-source fused initialization against best-single-source initialization. Large diamonds denote mean relative improvement across the 12 conditions; vertical whiskers show the frozen bootstrap 95% confidence interval.\n\n## 中文\n\n**图 X. 主要机制对比的配对相对改进（RMSE / sMAPE）。** 每个点表示一个 dataset–horizon 配对条件（D1/D4/D5/D6 × H1/H3/H5）的条件级相对改进。改进值均在对应消融臂之间成对计算，因此大于 0 表示所标示机制带来性能改善。C vs A 分离融合源信息初始化相对于原生初始化的贡献；C vs B 检验融合先验下可训练卷积骨干相对于冻结骨干的影响；D vs C 衡量 RevIN 的贡献；C vs E 比较多源融合初始化与最佳单源初始化。大菱形表示 12 个条件上的平均相对改进，纵向误差线表示冻结统计表中的 bootstrap 95% 置信区间。\n""",
        encoding="utf-8",
    )


def make_figure(metric: str, records: list[dict[str, object]], summaries: list[dict[str, object]], output_dir: Path):
    plt.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["DejaVu Serif"],
            "font.size": 9,
            "axes.titlesize": 12,
            "axes.titleweight": "bold",
            "axes.labelsize": 10,
            "xtick.labelsize": 8.5,
            "ytick.labelsize": 8.5,
            "legend.fontsize": 8,
            "legend.title_fontsize": 8,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "savefig.dpi": 600,
        }
    )
    figure_records = [row for row in records if row["metric"] == metric]
    figure_summaries = {row["contrast"]: row for row in summaries if row["metric"] == metric}
    jitter_values = np.linspace(-0.22, 0.22, len(CONDITION_ORDER))
    jitter = dict(zip(CONDITION_ORDER, jitter_values))

    all_values = [float(row["relative_improvement_pct"]) for row in figure_records]
    ci_values = []
    for row in figure_summaries.values():
        if row["bootstrap_95CI_low_pct"] is not None:
            ci_values.extend([float(row["bootstrap_95CI_low_pct"]), float(row["bootstrap_95CI_high_pct"])])
    y_min = min([0.0, *all_values, *ci_values])
    y_max = max([0.0, *all_values, *ci_values])
    span = max(y_max - y_min, 1.0)

    fig, ax = plt.subplots(figsize=(7.2, 4.65))
    for contrast_index, contrast in enumerate(CONTRASTS):
        subset = [row for row in figure_records if row["contrast"] == contrast["key"]]
        for row in subset:
            condition = (row["dataset"], int(row["horizon"]))
            ax.scatter(
                contrast_index + jitter[condition],
                float(row["relative_improvement_pct"]),
                s=48,
                marker=HORIZON_MARKERS[int(row["horizon"])],
                color=DATASET_COLORS[row["dataset"]],
                edgecolor="white",
                linewidth=0.65,
                alpha=0.95,
                zorder=3,
            )
        summary = figure_summaries[contrast["key"]]
        mean_value = float(summary["mean_RI_pct"])
        low = summary["bootstrap_95CI_low_pct"]
        high = summary["bootstrap_95CI_high_pct"]
        if low is not None and high is not None:
            ax.errorbar(
                contrast_index,
                mean_value,
                yerr=[[mean_value - float(low)], [float(high) - mean_value]],
                fmt="D",
                markersize=8,
                markerfacecolor="#2C3E50",
                markeredgecolor="white",
                markeredgewidth=0.75,
                ecolor="#2C3E50",
                elinewidth=1.8,
                capsize=4,
                capthick=1.2,
                zorder=5,
            )
        else:
            ax.scatter(
                contrast_index,
                mean_value,
                marker="D",
                s=64,
                color="#2C3E50",
                edgecolor="white",
                linewidth=0.75,
                zorder=5,
            )

    ax.axhline(0, color="#7F7F7F", linewidth=1.0, linestyle=(0, (4, 3)), zorder=1)
    ax.set_xlim(-0.55, len(CONTRASTS) - 0.45)
    ax.set_ylim(y_min - 0.10 * span, y_max + 0.15 * span)
    ax.set_xticks(range(len(CONTRASTS)))
    ax.set_xticklabels([contrast["label"] for contrast in CONTRASTS])
    ax.set_ylabel("Relative improvement (%)")
    ax.set_title(f"Paired Relative Improvements of Major Mechanism Contrasts ({metric})", pad=42)
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.7, alpha=0.65)
    ax.set_axisbelow(True)
    ax.text(
        0.99,
        0.985,
        "n = 12 paired conditions",
        transform=ax.transAxes,
        ha="right",
        va="top",
        fontsize=8,
        color="#555555",
    )
    ax.text(
        0.99,
        0.025,
        "mean marker; whisker = frozen bootstrap 95% CI",
        transform=ax.transAxes,
        ha="right",
        va="bottom",
        fontsize=7.5,
        color="#555555",
    )

    dataset_handles = [
        Line2D(
            [0],
            [0],
            marker="o",
            linestyle="None",
            markerfacecolor=DATASET_COLORS[dataset],
            markeredgecolor="white",
            markeredgewidth=0.6,
            markersize=6,
            label=dataset,
        )
        for dataset in EXPECTED_DATASETS
    ]
    horizon_handles = [
        Line2D(
            [0],
            [0],
            marker=HORIZON_MARKERS[horizon],
            linestyle="None",
            markerfacecolor="white",
            markeredgecolor="#333333",
            markersize=6,
            label=f"H{horizon}",
        )
        for horizon in EXPECTED_HORIZONS
    ]
    dataset_legend = ax.legend(
        handles=dataset_handles,
        title="Dataset",
        loc="upper left",
        bbox_to_anchor=(0.0, 1.21),
        ncol=4,
        frameon=False,
        handletextpad=0.35,
        columnspacing=0.8,
    )
    ax.add_artist(dataset_legend)
    ax.legend(
        handles=horizon_handles,
        title="Horizon",
        loc="upper right",
        bbox_to_anchor=(1.0, 1.21),
        ncol=3,
        frameon=False,
        handletextpad=0.35,
        columnspacing=0.8,
    )
    fig.subplots_adjust(left=0.105, right=0.985, bottom=0.17, top=0.78)

    stem = f"paired_relative_improvements_{metric.lower()}"
    for extension in ["png", "pdf", "svg"]:
        fig.savefig(output_dir / f"{stem}.{extension}", bbox_inches="tight", facecolor="white")
    plt.close(fig)


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Formal frozen A-E workbook")
    parser.add_argument("--output_dir", required=True, type=Path, help="New directory for figures and derived tables")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = args.input.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    if not input_path.is_file():
        raise FreezeValidationError(f"Input workbook does not exist: {input_path}")
    output_dir.mkdir(parents=True, exist_ok=True)
    sha256 = hashlib.sha256(input_path.read_bytes()).hexdigest()

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    means, source_sheet = load_condition_means(workbook)
    raw_audit = audit_raw_sheet(workbook)
    statistics = load_statistics(workbook)
    records_by_metric = derive_relative_improvements(means)
    all_records = records_by_metric["RMSE"] + records_by_metric["sMAPE"]
    summaries = []
    summaries.extend(summarize(records_by_metric["RMSE"], statistics, "RMSE"))
    summaries.extend(summarize(records_by_metric["sMAPE"], statistics, "sMAPE"))

    for row in summaries:
        frozen_mean = statistics.get((row["contrast"], row["metric"]), {}).get("frozen_mean")
        if frozen_mean is not None and abs(float(row["mean_RI_pct"]) - float(frozen_mean)) > 1e-8:
            raise FreezeValidationError(
                f"Derived mean disagrees with frozen statistic for {row['contrast']}/{row['metric']}: "
                f"derived={row['mean_RI_pct']}, frozen={frozen_mean}"
            )

    write_csv(output_dir / "paired_relative_improvements_rmse.csv", records_by_metric["RMSE"])
    write_csv(output_dir / "paired_relative_improvements_smape.csv", records_by_metric["sMAPE"])
    write_summary(output_dir / "paired_relative_improvements_summary.md", input_path, sha256, source_sheet, raw_audit, summaries)
    write_captions(output_dir / "paired_relative_improvements_captions.md")
    make_figure("RMSE", all_records, summaries, output_dir)
    make_figure("sMAPE", all_records, summaries, output_dir)

    missing_conditions = sorted(
        {
            (dataset, horizon)
            for dataset in EXPECTED_DATASETS
            for horizon in EXPECTED_HORIZONS
            if any((dataset, horizon, arm) not in means for arm in EXPECTED_ARMS)
        }
    )
    duplicate_conditions = []
    print(f"INPUT_WORKBOOK={input_path}")
    print(f"INPUT_SHA256={sha256}")
    print(f"SELECTED_DATA_SHEET={source_sheet}")
    print("TOTAL_CONDITIONS_EXPECTED=12")
    print(f"TOTAL_CONDITIONS_FOUND={len(set((row['dataset'], row['horizon']) for row in means.values()))}")
    print("CONTRASTS=4")
    print("METRICS=2")
    print("POINTS_PER_FIGURE_EXPECTED=48")
    print(f"POINTS_PER_FIGURE_FOUND={len(records_by_metric['RMSE'])}")
    print(f"MISSING_CONDITIONS={missing_conditions or 'NONE'}")
    print(f"DUPLICATE_CONDITIONS={duplicate_conditions or raw_audit.get('duplicates') or 'NONE'}")
    print(f"OUTPUT_DIR={output_dir}")
    print("GENERATED_FILES=10")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except FreezeValidationError as exc:
        print(f"VALIDATION_ERROR={exc}", file=sys.stderr)
        raise SystemExit(2)
