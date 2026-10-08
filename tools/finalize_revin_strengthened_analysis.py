#!/usr/bin/env python3
"""Audit and analyze completed RevIN-strengthened DLinear/TCN results.

This tool only reads the RevIN raw/prediction artifacts and the authoritative
freeze workbook. It writes analysis artifacts under the isolated robustness
directory and never writes formal baseline or freeze files.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import subprocess
import sys
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import openpyxl
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.baselines.revin_strengthened_protocol import (
    DEFAULT_FREEZE,
    EXPECTED_CELLS,
    PAPER_HORIZONS,
    PAPER_MODE,
    PAPER_SEEDS,
    PAPER_TARGETS,
    aggregate_five_seed,
    relative_error_change,
    relative_improvement,
    sha256_file,
)
from src.evaluation.metrics import compute_original_scale_metrics
from tools.build_final_statistical_derivation_freeze import (
    FORMAL_SHEETS,
    _normalize_main_rows,
    holm_adjust,
    read_sheet_rows,
    wilcoxon_exact,
)

RESULT_DIR = PROJECT_ROOT / "results" / "robustness" / "revin_strengthened_baselines"
RAW_PATH = RESULT_DIR / "REVIN_STRENGTHENED_BASELINES_5SEED_RAW.csv"
PREDICTIONS_PATH = RESULT_DIR / "REVIN_STRENGTHENED_BASELINES_PREDICTIONS.csv"
MANIFEST_PATH = RESULT_DIR / "REVIN_STRENGTHENED_BASELINES_MANIFEST.json"
AGG_PATH = RESULT_DIR / "REVIN_STRENGTHENED_BASELINES_5SEED_AGGREGATED.csv"
AUDIT_PATH = RESULT_DIR / "REVIN_STRENGTHENED_RESULTS_PROVENANCE_AUDIT.md"
XLSX_PATH = RESULT_DIR / "REVIN_STRENGTHENED_BASELINE_COMPARISON.xlsx"
REPORT_PATH = RESULT_DIR / "REVIN_STRENGTHENED_BASELINE_ROBUSTNESS_REPORT.md"

METHODS = ("DLinear_RevIN", "TCN_RevIN")
PARENTS = ("DLinear", "TCN")
KEY_COLUMNS = ["dataset", "target", "mode", "method", "horizon", "seed"]
CONDITIONS = [
    (dataset, target, horizon)
    for dataset, target in PAPER_TARGETS.items()
    for horizon in PAPER_HORIZONS
]
EXPECTED_KEYS = {
    (dataset, target, PAPER_MODE, method, horizon, seed)
    for dataset, target in PAPER_TARGETS.items()
    for method in METHODS
    for horizon in PAPER_HORIZONS
    for seed in PAPER_SEEDS
}


def _norm_frame(frame: pd.DataFrame) -> pd.DataFrame:
    output = frame.copy()
    if "dataset_id" in output.columns:
        output["dataset"] = output["dataset_id"]
    if "target_entity_key" in output.columns:
        output["target"] = output["target_entity_key"]
    if "scenario" in output.columns and "mode" not in output.columns:
        output["mode"] = output["scenario"]
    output["dataset"] = output["dataset"].astype(str).str.upper()
    output["target"] = output["target"].astype(str)
    output["mode"] = output["mode"].astype(str)
    output["method"] = output["method"].astype(str)
    output["horizon"] = output["horizon"].astype(int)
    output["seed"] = output["seed"].astype(int)
    return output


def _key_tuple(row: pd.Series) -> tuple[str, str, str, str, int, int]:
    return (
        str(row["dataset"]).upper(),
        str(row["target"]),
        str(row["mode"]),
        str(row["method"]),
        int(row["horizon"]),
        int(row["seed"]),
    )


def _keys(frame: pd.DataFrame) -> list[tuple[str, str, str, str, int, int]]:
    return [_key_tuple(row) for _, row in frame.iterrows()]


def _git_head() -> str:
    try:
        return subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=PROJECT_ROOT,
            check=True,
            capture_output=True,
            text=True,
        ).stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return "unavailable"


def _git_object_exists(commit: str) -> bool:
    if not commit:
        return False
    try:
        result = subprocess.run(
            ["git", "cat-file", "-e", f"{commit}^{{commit}}"],
            cwd=PROJECT_ROOT,
            check=False,
            capture_output=True,
            text=True,
        )
        return result.returncode == 0
    except OSError:
        return False


def _protected_snapshot() -> dict[str, str]:
    paths = [DEFAULT_FREEZE]
    formal_dir = PROJECT_ROOT / "outputs" / "baselines"
    if formal_dir.exists():
        paths.extend(sorted(path for path in formal_dir.rglob("*") if path.is_file()))
    snapshot: dict[str, str] = {}
    for path in paths:
        if path.exists() and path.is_file():
            snapshot[str(path.resolve())] = sha256_file(path)
    return snapshot


def _status(checks: list[dict[str, Any]], name: str, passed: bool, evidence: str, critical: bool = True) -> None:
    checks.append({"check": name, "status": "PASS" if passed else "FAIL", "critical": critical, "evidence": evidence})


def _load_formal() -> pd.DataFrame:
    main = _normalize_main_rows(read_sheet_rows(DEFAULT_FREEZE, FORMAL_SHEETS["main"]), FORMAL_SHEETS["main"])
    baseline = _normalize_main_rows(read_sheet_rows(DEFAULT_FREEZE, FORMAL_SHEETS["baseline"]), FORMAL_SHEETS["baseline"])
    frame = pd.DataFrame(main + baseline)
    frame = frame.loc[frame["method"].isin(("SITA",) + PARENTS)].copy()
    frame["dataset"] = frame["dataset"].astype(str).str.upper()
    frame["target"] = frame["target"].astype(str)
    frame["mode"] = "without"
    frame["horizon"] = frame["horizon"].astype(int)
    frame["seed"] = frame["seed"].astype(int)
    return frame


def _condition_aggregate(frame: pd.DataFrame) -> pd.DataFrame:
    rows = []
    group_cols = ["method", "dataset", "target", "horizon", "mode"]
    for identity, group in frame.groupby(group_cols, sort=True, dropna=False):
        method, dataset, target, horizon, mode = identity
        rows.append(
            {
                "method": method,
                "dataset": dataset,
                "target": target,
                "horizon": int(horizon),
                "mode": mode,
                "rmse_mean": float(group["rmse"].mean()),
                "rmse_std": float(group["rmse"].std(ddof=1)),
                "smape_mean": float(group["smape"].mean()),
                "smape_std": float(group["smape"].std(ddof=1)),
                "n_seeds": int(len(group)),
                "seeds": ",".join(str(seed) for seed in sorted(group["seed"].unique())),
            }
        )
    return pd.DataFrame(rows)


def _map_by_condition(frame: pd.DataFrame, method: str) -> dict[tuple[str, str, int], dict[str, float]]:
    subset = frame.loc[frame["method"].astype(str) == method].copy()
    output: dict[tuple[str, str, int], dict[str, float]] = {}
    for _, row in subset.iterrows():
        key = (str(row["dataset"]).upper(), str(row["target"]), int(row["horizon"]))
        output[key] = {
            "rmse": float(row["rmse_mean"] if "rmse_mean" in row else row["rmse"]),
            "smape": float(row["smape_mean"] if "smape_mean" in row else row["smape"]),
        }
    return output


def _write_if_changed(path: Path, content: str) -> None:
    if path.exists() and path.read_text(encoding="utf-8") == content:
        return
    path.write_text(content, encoding="utf-8")


def _audit() -> tuple[dict[str, Any], pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    checks: list[dict[str, Any]] = []
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    summary = _norm_frame(pd.read_csv(RAW_PATH))
    predictions = _norm_frame(pd.read_csv(PREDICTIONS_PATH))
    formal = _load_formal()

    _status(checks, "source files exist", all(path.exists() for path in (RAW_PATH, PREDICTIONS_PATH, MANIFEST_PATH, DEFAULT_FREEZE)),
            f"raw={RAW_PATH}; predictions={PREDICTIONS_PATH}; manifest={MANIFEST_PATH}; freeze={DEFAULT_FREEZE}")
    _status(checks, "manifest smoke flag", manifest.get("smoke") is False, f"smoke={manifest.get('smoke')!r}")
    _status(checks, "manifest code commit exists", _git_object_exists(str(manifest.get("code_commit", ""))),
            f"code_commit={manifest.get('code_commit')}")
    _status(checks, "manifest code commit matches current HEAD", str(manifest.get("code_commit")) == _git_head(),
            f"manifest={manifest.get('code_commit')}; HEAD={_git_head()}", critical=False)
    _status(checks, "run id is present and consistent", bool(manifest.get("run_id")) and summary["run_id"].nunique() == 1 and summary["run_id"].iloc[0] == manifest["run_id"],
            f"manifest_run_id={manifest.get('run_id')}; observed={summary['run_id'].dropna().unique().tolist()}")
    _status(checks, "manifest hashes match source artifacts",
            manifest.get("dlinear_sha256") == sha256_file(RESULT_DIR / "DLINEAR_REVIN_FORMAL_5SEED_RAW.csv")
            and manifest.get("tcn_sha256") == sha256_file(RESULT_DIR / "TCN_REVIN_FORMAL_5SEED_RAW.csv")
            and manifest.get("prediction_sha256") == sha256_file(PREDICTIONS_PATH),
            "dlinear/tcn/prediction SHA-256 matched manifest")

    observed_keys = set(_keys(summary))
    duplicate_count = int(sum(count > 1 for count in Counter(_keys(summary)).values()))
    _status(checks, "120-cell scope and identity completeness", len(observed_keys) == EXPECTED_CELLS and observed_keys == EXPECTED_KEYS,
            f"EXPECTED_CELLS={EXPECTED_CELLS}; ACTUAL_CELLS={len(observed_keys)}; missing={len(EXPECTED_KEYS - observed_keys)}; extra={len(observed_keys - EXPECTED_KEYS)}")
    _status(checks, "duplicate cells", duplicate_count == 0, f"DUPLICATE_CELLS={duplicate_count}")
    _status(checks, "method identities isolated", set(summary["method"]) == set(METHODS) and set(manifest.get("method_ids", [])) == {"BL6_DLinear_RevIN", "BL7_TCN_RevIN"},
            f"methods={sorted(summary['method'].unique())}; method_ids={manifest.get('method_ids')}")
    _status(checks, "dataset target horizon mode scope", set(summary["dataset"]) == set(PAPER_TARGETS) and all(set(summary.loc[summary["dataset"] == dataset, "target"]) == {target} for dataset, target in PAPER_TARGETS.items()) and set(summary["horizon"]) == set(PAPER_HORIZONS) and set(summary["mode"]) == {PAPER_MODE},
            f"targets={summary.groupby('dataset')['target'].unique().to_dict()}; horizons={sorted(summary['horizon'].unique())}; modes={sorted(summary['mode'].unique())}")
    _status(checks, "all seeds 42-46 present", set(summary["seed"]) == set(PAPER_SEEDS) and all(set(group["seed"]) == set(PAPER_SEEDS) for _, group in summary.groupby(["method", "dataset", "target", "horizon"])),
            f"seeds={sorted(summary['seed'].unique())}")
    _status(checks, "finite result metrics", bool(np.isfinite(summary[["rmse", "smape"]].to_numpy(dtype=float)).all()),
            f"rows={len(summary)}; finite_rmse_smape={bool(np.isfinite(summary[['rmse','smape']].to_numpy(dtype=float)).all())}")
    _status(checks, "runner provenance metadata", set(summary["result_status"]) == {"robustness_trial"} and set(summary["method_type"]) == {"robustness_baseline"} and set(summary["robustness_variant"]) == {"RevIN"},
            f"result_status={sorted(summary['result_status'].unique())}; method_type={sorted(summary['method_type'].unique())}; variant={sorted(summary['robustness_variant'].unique())}")
    _status(checks, "RevIN configuration provenance", bool((summary["revin_enabled"] == True).all()) and bool((summary["revin_affine"] == True).all()) and set(summary["revin_eps"].astype(float)) == {1e-5} and set(summary["revin_target_channel"].astype(int)) == {0},
            "enabled=True; affine=True; eps=1e-5; target_channel=0")
    _status(checks, "prediction and metric traceability fields", all(summary[column].notna().all() and summary[column].astype(str).str.len().gt(0).all() for column in ("sample_manifest_digest", "prediction_file_path", "code_commit", "protocol_version")),
            "sample_manifest_digest, prediction_file_path, code_commit, protocol_version are populated")
    _status(checks, "original sales metric space", set(summary["primary_metric_space"].astype(str)) == {"original_sales"} and all(set(summary[column].astype(str)) == {"original_sales_space"} for column in ("metric_space", "rmse_metric_space", "smape_metric_space")) and set(predictions["metric_space"].astype(str)) == {"original_sales_space"},
            "primary_metric_space=original_sales; summary metric_space/rmse_metric_space/smape_metric_space and predictions=original_sales_space")

    placeholder_hits = []
    for path in (RAW_PATH, PREDICTIONS_PATH, MANIFEST_PATH):
        text = path.read_text(encoding="utf-8").lower()
        for token in ("placeholder", "mock", "synthetic", "fake", "dummy", "random_fixture"):
            if token in text:
                placeholder_hits.append(f"{path.name}:{token}")
    _status(checks, "no placeholder/mock/synthetic markers", not placeholder_hits, f"hits={placeholder_hits}")

    prediction_keys = set(_keys(predictions))
    sample_key_duplicate_groups = []
    for key, group in predictions.groupby(KEY_COLUMNS, sort=False):
        if group["sample_key"].duplicated().any():
            sample_key_duplicate_groups.append(key)
    _status(checks, "prediction groups cover all 120 cells", prediction_keys == EXPECTED_KEYS, f"groups={len(prediction_keys)}; expected={len(EXPECTED_KEYS)}")
    _status(checks, "prediction sample keys are unique within cell", not sample_key_duplicate_groups and predictions["sample_key"].notna().all(), f"duplicate_sample_key_groups={len(sample_key_duplicate_groups)}")
    recompute_failures = []
    sample_count_failures = []
    for key, group in predictions.groupby(KEY_COLUMNS, sort=False):
        tuple_key = tuple(key[:3]) + (key[3], int(key[4]), int(key[5]))
        source_row = summary.loc[summary[KEY_COLUMNS].apply(tuple, axis=1) == tuple_key]
        if len(source_row) != 1:
            recompute_failures.append(f"missing summary for {tuple_key}")
            continue
        expected = source_row.iloc[0]
        if len(group) != int(expected["sample_count"]):
            sample_count_failures.append(f"{tuple_key}:pred={len(group)} summary={int(expected['sample_count'])}")
        metrics = compute_original_scale_metrics(group["y_true"].to_numpy(float), group["y_pred"].to_numpy(float))
        if not math.isclose(metrics["rmse"], float(expected["rmse"]), rel_tol=1e-10, abs_tol=1e-10) or not math.isclose(metrics["smape"], float(expected["smape"]), rel_tol=1e-10, abs_tol=1e-10):
            recompute_failures.append(f"{tuple_key}:computed=({metrics['rmse']},{metrics['smape']}) summary=({expected['rmse']},{expected['smape']})")
    _status(checks, "prediction rows reproduce summary metrics", not recompute_failures, f"mismatches={len(recompute_failures)}")
    _status(checks, "prediction counts match summary sample_count", not sample_count_failures, f"mismatches={len(sample_count_failures)}")
    _status(checks, "prediction values are finite", bool(np.isfinite(predictions[["y_true", "y_pred"]].to_numpy(float)).all()), f"prediction_rows={len(predictions)}")

    formal_counts = formal.groupby("method").size().to_dict()
    formal_keys = formal.groupby(["method", "dataset", "target", "horizon", "seed"]).size()
    _status(checks, "formal SITA/DLinear/TCN scope loaded read-only", formal_counts == {"SITA": 60, "DLinear": 60, "TCN": 60} and bool((formal_keys == 1).all()), f"counts={formal_counts}; duplicate_or_missing_keys={int((formal_keys != 1).sum())}")
    _status(checks, "formal scope matches requested authority", set(formal["dataset"]) == set(PAPER_TARGETS) and set(formal["horizon"]) == set(PAPER_HORIZONS) and set(formal["seed"]) == set(PAPER_SEEDS) and set(formal["metric_space"]) == {"original_sales_space"}, f"freeze={DEFAULT_FREEZE}; metric_space={sorted(formal['metric_space'].unique())}")

    formal_parent_maps = {method: _map_by_condition(_condition_aggregate(formal), method) for method in PARENTS}
    rev_maps = {method: _map_by_condition(aggregate_five_seed(summary.loc[summary["method"] == method]), method) for method in METHODS}
    identity_differences = {}
    for parent, strengthened in (("DLinear", "DLinear_RevIN"), ("TCN", "TCN_RevIN")):
        rmse_diff = 0
        smape_diff = 0
        for dataset, target, horizon in CONDITIONS:
            parent_value = formal_parent_maps[parent][(dataset, target, horizon)]
            rev_value = rev_maps[strengthened][(dataset, target, horizon)]
            rmse_diff += int(not math.isclose(parent_value["rmse"], rev_value["rmse"], rel_tol=0.0, abs_tol=0.0))
            smape_diff += int(not math.isclose(parent_value["smape"], rev_value["smape"], rel_tol=0.0, abs_tol=0.0))
        identity_differences[parent] = {"rmse_condition_differences": rmse_diff, "smape_condition_differences": smape_diff}
    _status(checks, "DLinear_RevIN is not a complete copy of DLinear", identity_differences["DLinear"]["rmse_condition_differences"] > 0 and identity_differences["DLinear"]["smape_condition_differences"] > 0, str(identity_differences["DLinear"]))
    _status(checks, "TCN_RevIN is not a complete copy of TCN", identity_differences["TCN"]["rmse_condition_differences"] > 0 and identity_differences["TCN"]["smape_condition_differences"] > 0, str(identity_differences["TCN"]))

    audit = {
        "checks": checks,
        "manifest": manifest,
        "summary_rows": int(len(summary)),
        "prediction_rows": int(len(predictions)),
        "formal_rows": int(len(formal)),
        "identity_differences": identity_differences,
        "raw_sha256": sha256_file(RAW_PATH),
        "predictions_sha256": sha256_file(PREDICTIONS_PATH),
        "manifest_sha256": sha256_file(MANIFEST_PATH),
        "freeze_sha256_before": sha256_file(DEFAULT_FREEZE),
    }
    return audit, summary, predictions, formal


def _audit_markdown(audit: dict[str, Any]) -> str:
    checks = audit["checks"]
    failures = [row for row in checks if row["status"] == "FAIL" and row["critical"]]
    lines = [
        "# RevIN-Strengthened Results Provenance Audit",
        "",
        f"Audit date: 2026-09-12  ",
        f"Repository: `{PROJECT_ROOT}`  ",
        f"Result directory: `{RESULT_DIR}`  ",
        "",
        f"`REVIN_STRENGTHENED_RESULTS_PROVENANCE_AUDIT={'FAIL' if failures else 'PASS'}`",
        f"`CRITICAL_FAILURES={len(failures)}`",
        "",
        "## Scope and authority",
        "",
        f"- RevIN source raw: `{RAW_PATH}` (SHA-256 `{audit['raw_sha256']}`).",
        f"- Prediction source: `{PREDICTIONS_PATH}` (SHA-256 `{audit['predictions_sha256']}`).",
        f"- Manifest: `{MANIFEST_PATH}` (SHA-256 `{audit['manifest_sha256']}`).",
        f"- Formal read-only authority: `{DEFAULT_FREEZE}` (SHA-256 before analysis `{audit['freeze_sha256_before']}`).",
        "- The RevIN rows retain `result_status=robustness_trial`; they are independent robustness-run outputs, not promoted formal freeze rows.",
        "- No paper conclusion is permitted if a critical provenance check fails.",
        "",
        "## Checks",
        "",
        "| Check | Status | Evidence |",
        "|---|---|---|",
    ]
    for row in checks:
        evidence = str(row["evidence"]).replace("|", "\\|").replace("\n", " ")
        lines.append(f"| {row['check']} | {row['status']} | {evidence} |")
    lines.extend([
        "",
        "## Interpretation boundary",
        "",
        "The audit verifies that each method × dataset × horizon × seed identity is unique, the prediction groups are traceable and reproduce the stored original-sales RMSE/sMAPE, and the strengthened methods are not numerically identical to their frozen parent methods. It does not retroactively promote these robustness-trial rows into the formal freeze.",
        "",
    ])
    return "\n".join(lines)


def _comparison_tables(summary: pd.DataFrame, formal: pd.DataFrame) -> dict[str, pd.DataFrame]:
    rev_aggregated = aggregate_five_seed(summary)
    rev_aggregated = rev_aggregated.sort_values(["method", "dataset", "horizon"]).reset_index(drop=True)
    agg_path_text = rev_aggregated.to_csv(index=False, encoding="utf-8")
    _write_if_changed(AGG_PATH, agg_path_text)

    formal_agg = _condition_aggregate(formal)
    sita_map = _map_by_condition(formal_agg, "SITA")
    effect_rows = []
    versus_by_method: dict[str, list[dict[str, Any]]] = {"DLinear_RevIN": [], "TCN_RevIN": []}
    parent_for = {"DLinear_RevIN": "DLinear", "TCN_RevIN": "TCN"}
    rev_maps = {method: _map_by_condition(rev_aggregated, method) for method in METHODS}
    formal_maps = {method: _map_by_condition(formal_agg, method) for method in ("SITA",) + PARENTS}

    for strengthened, parent in parent_for.items():
        for dataset, target, horizon in CONDITIONS:
            key = (dataset, target, horizon)
            parent_values = formal_maps[parent][key]
            strengthened_values = rev_maps[strengthened][key]
            for metric, label in (("rmse", "RMSE"), ("smape", "sMAPE")):
                before = parent_values[metric]
                after = strengthened_values[metric]
                effect_rows.append({
                    "dataset": dataset,
                    "target": target,
                    "horizon": horizon,
                    "mode": PAPER_MODE,
                    "parent_method": parent,
                    "strengthened_method": strengthened,
                    "metric": label,
                    "parent_error": before,
                    "strengthened_error": after,
                    "absolute_change_strengthened_minus_parent": after - before,
                    "relative_change_pct_strengthened_minus_parent": relative_error_change(before, after),
                    "interpretation": "improved" if after < before else "degraded" if after > before else "tie",
                })
            for metric, label in (("rmse", "RMSE"), ("smape", "sMAPE")):
                sita = sita_map[key][metric]
                baseline = strengthened_values[metric]
                improvement = relative_improvement(baseline, sita)
                versus_by_method[strengthened].append({
                    "dataset": dataset,
                    "target": target,
                    "horizon": horizon,
                    "mode": PAPER_MODE,
                    "metric": label,
                    "sita_error": sita,
                    "strengthened_method": strengthened,
                    "strengthened_error": baseline,
                    "relative_improvement_pct": improvement,
                    "outcome": "WIN" if sita < baseline else "LOSS" if sita > baseline else "TIE",
                    "effect_direction": "SITA lower error" if sita < baseline else "strengthened baseline lower error" if sita > baseline else "exact tie",
                })

    effect = pd.DataFrame(effect_rows)
    versus_frames = {method: pd.DataFrame(rows) for method, rows in versus_by_method.items()}
    stats_rows = []
    raw_p: dict[str, float] = {}
    for strengthened, rows in versus_frames.items():
        parent = parent_for[strengthened]
        for metric in ("RMSE", "sMAPE"):
            subset = rows.loc[rows["metric"] == metric].copy()
            differences = (subset["sita_error"] - subset["strengthened_error"]).astype(float).tolist()
            test = wilcoxon_exact(differences)
            key = f"SITA_vs_{strengthened}|{metric}"
            raw_p[key] = float(test["p_value"])
            improvements = subset["relative_improvement_pct"].astype(float).tolist()
            wins = int((subset["outcome"] == "WIN").sum())
            losses = int((subset["outcome"] == "LOSS").sum())
            ties = int((subset["outcome"] == "TIE").sum())
            mean_difference = float(np.mean(differences))
            stats_rows.append({
                "comparison": f"SITA_vs_{strengthened}",
                "parent_method": parent,
                "strengthened_method": strengthened,
                "metric": metric,
                "sample_size": len(differences),
                "sita_wins": wins,
                "sita_losses": losses,
                "sita_ties": ties,
                "mean_paired_difference_sita_minus_strengthened": mean_difference,
                "median_paired_difference_sita_minus_strengthened": float(np.median(differences)),
                "effect_direction": "SITA lower error overall" if mean_difference < 0 else "strengthened baseline lower error overall" if mean_difference > 0 else "no mean direction",
                "wilcoxon_statistic_W": test["statistic"],
                "wilcoxon_w_plus": test["w_plus"],
                "wilcoxon_w_minus": test["w_minus"],
                "wilcoxon_raw_p": test["p_value"],
                "wilcoxon_alternative": test["alternative"],
                "wilcoxon_method": test["method"],
                "zero_difference_count": test["zero_difference_count"],
                "tied_absolute_difference_count": test["tied_absolute_difference_count"],
                "mean_relative_improvement_pct": float(np.mean(improvements)),
                "median_relative_improvement_pct": float(np.median(improvements)),
                "min_relative_improvement_pct": float(np.min(improvements)),
                "max_relative_improvement_pct": float(np.max(improvements)),
                "statistical_unit": "12 dataset-horizon condition means",
                "tie_rule": "exact numeric no tolerance",
            })
    adjusted = holm_adjust(raw_p)
    for row in stats_rows:
        key = f"{row['comparison'].replace('SITA_vs_', 'SITA_vs_')}|{row['metric']}"
        row["holm_adjusted_p"] = adjusted[key]["adjusted_p"]
        row["holm_alpha"] = 0.05
        row["holm_significant_0.05"] = bool(row["holm_adjusted_p"] < 0.05)
        row["holm_family_size"] = len(raw_p)
    stats = pd.DataFrame(stats_rows)

    summary_rows = []
    for strengthened, frame in versus_frames.items():
        for metric in ("RMSE", "sMAPE"):
            subset = frame.loc[frame["metric"] == metric]
            summary_rows.append({
                "comparison": f"SITA_vs_{strengthened}",
                "strengthened_method": strengthened,
                "metric": metric,
                "n_conditions": int(len(subset)),
                "wins": int((subset["outcome"] == "WIN").sum()),
                "losses": int((subset["outcome"] == "LOSS").sum()),
                "ties": int((subset["outcome"] == "TIE").sum()),
                "mean_relative_improvement_pct": float(subset["relative_improvement_pct"].mean()),
                "median_relative_improvement_pct": float(subset["relative_improvement_pct"].median()),
                "min_relative_improvement_pct": float(subset["relative_improvement_pct"].min()),
                "max_relative_improvement_pct": float(subset["relative_improvement_pct"].max()),
            })
    relative = pd.DataFrame(summary_rows)

    win_loss_rows = []
    for row in summary_rows:
        win_loss_rows.append({
            "comparison": row["comparison"],
            "metric": row["metric"],
            "n_conditions": row["n_conditions"],
            "wins": row["wins"],
            "losses": row["losses"],
            "ties": row["ties"],
            "win_cells": "; ".join(f"{r['dataset']} H{r['horizon']}" for _, r in versus_frames[row["strengthened_method"]].loc[(versus_frames[row["strengthened_method"]]["metric"] == row["metric"]) & (versus_frames[row["strengthened_method"]]["outcome"] == "WIN")].iterrows()),
            "loss_cells": "; ".join(f"{r['dataset']} H{r['horizon']}" for _, r in versus_frames[row["strengthened_method"]].loc[(versus_frames[row["strengthened_method"]]["metric"] == row["metric"]) & (versus_frames[row["strengthened_method"]]["outcome"] == "LOSS")].iterrows()),
            "tie_cells": "; ".join(f"{r['dataset']} H{r['horizon']}" for _, r in versus_frames[row["strengthened_method"]].loc[(versus_frames[row["strengthened_method"]]["metric"] == row["metric"]) & (versus_frames[row["strengthened_method"]]["outcome"] == "TIE")].iterrows()),
        })
    win_loss = pd.DataFrame(win_loss_rows)

    return {
        "aggregated": rev_aggregated,
        "formal_aggregated": formal_agg,
        "effect": effect,
        "versus": versus_frames,
        "stats": stats,
        "relative": relative,
        "win_loss": win_loss,
    }


def _cells_text(frame: pd.DataFrame, outcome: str) -> str:
    rows = frame.loc[frame["outcome"] == outcome]
    return ", ".join(f"{row.dataset} H{int(row.horizon)}" for row in rows.itertuples()) or "none"


def _verdict(relative: pd.DataFrame) -> str:
    means_positive = bool((relative["mean_relative_improvement_pct"] > 0).all())
    wins_outnumber = bool((relative["wins"] > relative["losses"]).all())
    any_losses = bool((relative["losses"] > 0).any())
    if means_positive and wins_outnumber:
        return "MIXED" if any_losses else "SUPPORTED"
    return "NOT_SUPPORTED"


def _report(audit: dict[str, Any], tables: dict[str, pd.DataFrame], freeze_sha_after: str) -> str:
    relative = tables["relative"]
    stats = tables["stats"]
    verdict = _verdict(relative)
    lines = [
        "# RevIN-Strengthened Baseline Robustness Report",
        "",
        "Scope: completed DLinear_RevIN and TCN_RevIN robustness results only; no retraining was performed.",
        "",
        f"`FINAL_ROBUSTNESS_VERDICT={verdict}`",
        f"`PROVENANCE_AUDIT=PASS`",
        "",
        "## Matrix and protection boundary",
        "",
        "- Matrix: D1 `1/10`, D4 `166/258`, D5 `48/1159414`, D6 `CA_1/FOODS_3_080`; H=1,3,5; mode=`without`; seeds=42–46.",
        "- RevIN strengthened results: 120/120 unique cells; five-seed aggregation: 24 rows.",
        f"- Formal read-only freeze: `{DEFAULT_FREEZE}`; SHA-256 after analysis `{freeze_sha_after}`.",
        "- Freeze workbook, original formal baseline artifacts, SITA formal results, and paper master tables were not written.",
        "",
        "## RevIN effect on the parent baseline",
        "",
        "Negative `change` means the strengthened baseline error decreased; relative change is `(strengthened-parent)/parent*100`.",
        "",
        "| Method | Metric | Mean absolute change | Mean relative change (%) | Improved cells | Degraded cells |",
        "|---|---:|---:|---:|---:|---:|",
    ]
    for method, parent in (("DLinear_RevIN", "DLinear"), ("TCN_RevIN", "TCN")):
        effect = tables["effect"].loc[tables["effect"]["strengthened_method"] == method]
        for metric in ("RMSE", "sMAPE"):
            subset = effect.loc[effect["metric"] == metric]
            lines.append(f"| {method} vs {parent} | {metric} | {subset['absolute_change_strengthened_minus_parent'].mean():.6f} | {subset['relative_change_pct_strengthened_minus_parent'].mean():.3f} | {int((subset['interpretation'] == 'improved').sum())} | {int((subset['interpretation'] == 'degraded').sum())} |")
    lines.extend(["", "## SITA versus strengthened baselines", "", "Positive relative improvement means SITA has lower error.", "", "| Comparison | Metric | Wins | Losses | Ties | Mean improvement (%) | Median improvement (%) | Min (%) | Max (%) |", "|---|---:|---:|---:|---:|---:|---:|---:|---:|"])
    for row in relative.itertuples():
        lines.append(f"| {row.comparison} | {row.metric} | {row.wins} | {row.losses} | {row.ties} | {row.mean_relative_improvement_pct:.3f} | {row.median_relative_improvement_pct:.3f} | {row.min_relative_improvement_pct:.3f} | {row.max_relative_improvement_pct:.3f} |")
    lines.extend(["", "## Concrete cells", "", "The following lists are condition-level five-seed means; `SITA < strengthened` is a SITA win because lower error is better.", ""])
    for method, frame in tables["versus"].items():
        lines.extend([f"### SITA versus {method}", ""])
        for metric in ("RMSE", "sMAPE"):
            subset = frame.loc[frame["metric"] == metric]
            lines.append(f"- {metric} SITA < {method}: {_cells_text(subset, 'WIN')}")
            lines.append(f"- {metric} SITA > {method} (SITA losses): {_cells_text(subset, 'LOSS')}")
            lines.append(f"- {metric} exact ties: {_cells_text(subset, 'TIE')}")
        lines.append("")
    lines.extend(["## Wilcoxon signed-rank and Holm correction", "", "Paired unit is the same 12 dataset × horizon condition means. Differences are `SITA - strengthened`; the existing exact two-sided Wilcoxon implementation and Holm step-down correction are reused. No seed or condition was changed for significance testing.", "", "| Comparison | Metric | n | W | W+ | W- | Raw p | Holm adjusted p | Direction | Significant (0.05) |", "|---|---:|---:|---:|---:|---:|---:|---:|---|---|"])
    for _, row in stats.iterrows():
        lines.append(f"| {row['comparison']} | {row['metric']} | {int(row['sample_size'])} | {row['wilcoxon_statistic_W']:.3f} | {row['wilcoxon_w_plus']:.3f} | {row['wilcoxon_w_minus']:.3f} | {row['wilcoxon_raw_p']:.6g} | {row['holm_adjusted_p']:.6g} | {row['effect_direction']} | {'YES' if bool(row['holm_significant_0.05']) else 'NO'} |")
    lines.extend(["", "## Robustness judgment", "", "The verdict rule used here is explicit: `SUPPORTED` requires positive mean improvement and more SITA wins than losses for all four comparisons, with no losses; `MIXED` uses the same overall direction but has at least one strengthened-baseline win; `NOT_SUPPORTED` is used otherwise.", "", f"`FINAL_ROBUSTNESS_VERDICT={verdict}`", ""])
    if verdict == "SUPPORTED":
        lines.append("SITA retains an overall advantage after strengthening representative deep-learning baselines with RevIN. Therefore, the observed SITA advantage cannot be explained solely by the use of RevIN.")
    elif verdict == "MIXED":
        lines.append("RevIN narrows the performance gap in several conditions, while SITA retains an overall advantage across the evaluated settings.")
    else:
        lines.append("The strengthened baselines remove most or all of the observed SITA advantage, so the existing evidence is insufficient to exclude RevIN as a major source of the performance gain.")
    lines.extend(["", "No additive contribution percentage is claimed. RevIN effects are reported as paired error changes only; ablation deltas are not treated as additive contribution shares.", ""])
    return "\n".join(lines)


def _style_sheet(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.font = Font(name="Arial", size=10, color="000000")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        max_len = max(len(str(cell.value or "")) for cell in column_cells[:80])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)


def _write_df(ws: openpyxl.worksheet.worksheet.Worksheet, frame: pd.DataFrame) -> None:
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        ws.append(list(row))
    _style_sheet(ws)


def _write_xlsx(tables: dict[str, pd.DataFrame], audit: dict[str, Any], freeze_sha_after: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    readme = wb.create_sheet("README")
    readme.append(["Item", "Value"])
    for item, value in [
        ("Scope", "Completed RevIN-strengthened DLinear/TCN analysis; no retraining"),
        ("Matrix", "D1/D4/D5/D6 × H1/H3/H5 × seeds 42-46 × without"),
        ("Metric space", "original_sales_space; sMAPE is percent"),
        ("Primary authority", str(DEFAULT_FREEZE)),
        ("Freeze SHA-256 after analysis", freeze_sha_after),
        ("Provenance audit", "PASS"),
        ("Final robustness verdict", _verdict(tables["relative"])),
        ("Wilcoxon", "paired exact two-sided on 12 condition means"),
        ("Holm", "step-down correction over four tests: two methods × two metrics"),
        ("Contribution boundary", "No additive RevIN/source-initialization contribution percentages claimed"),
    ]:
        readme.append([item, value])
    _style_sheet(readme)

    _write_df(wb.create_sheet("Aggregated"), tables["aggregated"])
    for method, sheet_name in (("DLinear_RevIN", "RevIN_Effect_DLinear"), ("TCN_RevIN", "RevIN_Effect_TCN")):
        source = tables["effect"].loc[tables["effect"]["strengthened_method"] == method].copy()
        source = source.sort_values(["metric", "dataset", "horizon"]).reset_index(drop=True)
        ws = wb.create_sheet(sheet_name)
        headers = ["dataset", "target", "horizon", "mode", "parent_method", "strengthened_method", "metric", "parent_error", "strengthened_error", "absolute_change", "relative_change_pct", "interpretation"]
        ws.append(headers)
        for idx, row in source.iterrows():
            excel_row = idx + 2
            ws.append([row["dataset"], row["target"], int(row["horizon"]), row["mode"], row["parent_method"], row["strengthened_method"], row["metric"], row["parent_error"], row["strengthened_error"], f"=I{excel_row}-H{excel_row}", f"=IFERROR(J{excel_row}/H{excel_row}*100,\"\")", f"=IF(J{excel_row}<0,\"improved\",IF(J{excel_row}>0,\"degraded\",\"tie\"))"])
        _style_sheet(ws)

    for method, sheet_name in (("DLinear_RevIN", "SITA_vs_DLinear_RevIN"), ("TCN_RevIN", "SITA_vs_TCN_RevIN")):
        source = tables["versus"][method].sort_values(["metric", "dataset", "horizon"]).reset_index(drop=True)
        ws = wb.create_sheet(sheet_name)
        headers = ["dataset", "target", "horizon", "mode", "metric", "SITA_error", "strengthened_method", "strengthened_error", "relative_improvement_pct", "outcome", "effect_direction"]
        ws.append(headers)
        for idx, row in source.iterrows():
            excel_row = idx + 2
            ws.append([row["dataset"], row["target"], int(row["horizon"]), row["mode"], row["metric"], row["sita_error"], row["strengthened_method"], row["strengthened_error"], f"=IFERROR((H{excel_row}-F{excel_row})/H{excel_row}*100,\"\")", f"=IF(F{excel_row}<H{excel_row},\"WIN\",IF(F{excel_row}>H{excel_row},\"LOSS\",\"TIE\"))", f"=IF(F{excel_row}<H{excel_row},\"SITA lower error\",IF(F{excel_row}>H{excel_row},\"strengthened baseline lower error\",\"exact tie\"))"])
        _style_sheet(ws)

    wlt = wb.create_sheet("Win_Loss_Tie")
    wlt_headers = ["comparison", "metric", "n_conditions", "wins", "losses", "ties", "win_cells", "loss_cells", "tie_cells"]
    wlt.append(wlt_headers)
    for idx, row in tables["win_loss"].iterrows():
        sheet = "SITA_vs_DLinear_RevIN" if row["comparison"].endswith("DLinear_RevIN") else "SITA_vs_TCN_RevIN"
        excel_row = idx + 2
        wlt.append([row["comparison"], row["metric"], f'=COUNTIF(\'{sheet}\'!$E$2:$E$25,B{excel_row})', f'=COUNTIFS(\'{sheet}\'!$E$2:$E$25,B{excel_row},\'{sheet}\'!$J$2:$J$25,"WIN")', f'=COUNTIFS(\'{sheet}\'!$E$2:$E$25,B{excel_row},\'{sheet}\'!$J$2:$J$25,"LOSS")', f'=COUNTIFS(\'{sheet}\'!$E$2:$E$25,B{excel_row},\'{sheet}\'!$J$2:$J$25,"TIE")', row["win_cells"], row["loss_cells"], row["tie_cells"]])
    _style_sheet(wlt)

    rel = wb.create_sheet("Relative_Improvement")
    rel_headers = ["comparison", "metric", "n_conditions", "mean_relative_improvement_pct", "median_relative_improvement_pct", "min_relative_improvement_pct", "max_relative_improvement_pct"]
    rel.append(rel_headers)
    for idx, row in tables["relative"].iterrows():
        sheet = "SITA_vs_DLinear_RevIN" if row["comparison"].endswith("DLinear_RevIN") else "SITA_vs_TCN_RevIN"
        if row["metric"] == "RMSE":
            start, end = 2, 13
        else:
            start, end = 14, 25
        excel_row = idx + 2
        rel.append([row["comparison"], row["metric"], f'=COUNTIF(\'{sheet}\'!$E$2:$E$25,B{excel_row})', f'=AVERAGE(\'{sheet}\'!$I${start}:$I${end})', f'=MEDIAN(\'{sheet}\'!$I${start}:$I${end})', f'=MIN(\'{sheet}\'!$I${start}:$I${end})', f'=MAX(\'{sheet}\'!$I${start}:$I${end})'])
    _style_sheet(rel)

    _write_df(wb.create_sheet("Statistical_Tests"), tables["stats"])
    raw = pd.read_csv(RAW_PATH)
    trace = pd.DataFrame([
        {"artifact": "RevIN raw summary", "path": str(RAW_PATH), "sha256": sha256_file(RAW_PATH), "role": "120 seed-level result rows"},
        {"artifact": "RevIN predictions", "path": str(PREDICTIONS_PATH), "sha256": sha256_file(PREDICTIONS_PATH), "role": "prediction and y_true traceability"},
        {"artifact": "RevIN manifest", "path": str(MANIFEST_PATH), "sha256": sha256_file(MANIFEST_PATH), "role": "run/protocol/code provenance"},
        {"artifact": "Formal freeze authority", "path": str(DEFAULT_FREEZE), "sha256": freeze_sha_after, "role": "read-only SITA/DLinear/TCN source"},
        {"artifact": "Analysis code commit", "path": "git HEAD", "sha256": _git_head(), "role": "analysis provenance"},
        {"artifact": "Raw summary rows", "path": "in-memory", "sha256": str(len(raw)), "role": "120 expected cells"},
    ])
    _write_df(wb.create_sheet("Traceability"), trace)
    _write_df(wb.create_sheet("Frozen_Official"), tables["formal_aggregated"])
    _write_df(wb.create_sheet("RevIN_Raw"), _norm_frame(pd.read_csv(RAW_PATH)))
    wb.save(XLSX_PATH)


def run(audit_only: bool = False) -> None:
    if not RESULT_DIR.exists():
        raise FileNotFoundError(RESULT_DIR)
    protected_before = _protected_snapshot()
    audit, summary, predictions, formal = _audit()
    _write_if_changed(AUDIT_PATH, _audit_markdown(audit))
    failures = [row for row in audit["checks"] if row["status"] == "FAIL" and row["critical"]]
    if failures:
        raise RuntimeError(f"provenance audit failed; see {AUDIT_PATH}")
    if audit_only:
        print("REVIN_STRENGTHENED_RESULTS_PROVENANCE_AUDIT=PASS")
        print(f"AUDIT_PATH={AUDIT_PATH}")
        return

    tables = _comparison_tables(summary, formal)
    freeze_sha_after = sha256_file(DEFAULT_FREEZE)
    if freeze_sha_after != audit["freeze_sha256_before"]:
        raise RuntimeError("formal freeze hash changed during read-only comparison")
    _write_if_changed(REPORT_PATH, _report(audit, tables, freeze_sha_after))
    _write_xlsx(tables, audit, freeze_sha_after)
    protected_after = _protected_snapshot()
    if protected_before != protected_after:
        raise RuntimeError("protected formal files changed during analysis")
    print(f"REVIN_STRENGTHENED_RESULTS_PROVENANCE_AUDIT=PASS")
    print(f"EXPECTED_CELLS={EXPECTED_CELLS}")
    print(f"ACTUAL_CELLS={len(summary)}")
    print(f"FINAL_ROBUSTNESS_VERDICT={_verdict(tables['relative'])}")
    print(f"AUDIT_PATH={AUDIT_PATH}")
    print(f"AGG_PATH={AGG_PATH}")
    print(f"XLSX_PATH={XLSX_PATH}")
    print(f"REPORT_PATH={REPORT_PATH}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit-only", action="store_true")
    args = parser.parse_args()
    run(audit_only=args.audit_only)


if __name__ == "__main__":
    main()
