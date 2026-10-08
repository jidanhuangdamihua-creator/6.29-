#!/usr/bin/env python3
"""Build D4/D5/D6 70-target workbooks: new 60-target BL+SITA plus old 10-target BL+SITA."""

from __future__ import annotations

import ast
import csv
import json
import math
import statistics
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

import argparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OLD_ROOT = ROOT / "outputs" / "additional_target_robustness_20261004"
OLD_FILL_ROOT = OLD_ROOT / "old10_missing_baselines"
PROTOCOL = "405beca0c84a85d3bdd77016cda37d5c2744ca7d4bc863e5c9d77ab17cd7788c"
OLD_MANIFEST_SHA = "c83d03d337ef2b8c31d1e6b8aa0ea94f2b4ae8f7c6b312c39aedfdaf4035fec9"
DATASETS = {
    "D4": {
        "dataset": "D4",
        "store": "155, 274, 576, 166, 139, 161",
        "stores": ["155", "274", "576", "166", "139", "161"],
        "group_field": "same_store_only",
        "group_value": "True",
        "new_root": ROOT / "outputs" / "d4_multistore_60target_robustness_20261004",
        "manifest_name": "d4_60target_manifest.json",
        "out_name": "D4_70target_BL_SITA.xlsx",
        "new_manifest_sha256": "1a4ad685b60563d3e778c3d656535fa5c101e4cd1c041c0dc9b4e569468fea84",
        "source_sha256": "17a1fa5bd1dddfd46bda2a6922ff7821aee2a7e79deca58a94ff7bf20821f7ef",
        "universe_id": (
            "155=203f689fc16e9fa2; 274=be38811222404be3; 576=d232830006b2824f; "
            "166=862ab74463fdbbe1; 139=e18825d43332dc3f; 161=3f0fbe7b41148e89"
        ),
        "old_source_data_path": "数据集/固化数据/d1_d6_sealed_v1/dataset4/source.parquet",
        "protocol_version": "d4_multistore_60target_v1",
        "old_source_universe_id": "c407fe5b6f169ea2",
        "identity_overlap_note": "old additional 10 are store 166; those 10 identities also appear in the new 60 (different source_universe_id)",
    },
    "D5": {
        "dataset": "D5",
        "store": "48",
        "group_field": "family",
        "group_value": "GROCERY I",
        "new_root": ROOT / "outputs" / "d5_60target_robustness_20261004",
        "manifest_name": "d5_60target_manifest.json",
        "out_name": "D5_70target_BL_SITA.xlsx",
        "new_manifest_sha256": "613f992d448a2717f6ab119bd10a3d2ebc7a708e72aca28f35020148f86c63b8",
        "source_sha256": "0cbbda5ae42e2e962e7fd83a4b57a10d4fc0fdafc2687177bee0841e6d40e506",
        "universe_id": "755448b5dc59eb89",
        "old_source_data_path": "数据集/固化数据/d1_d6_sealed_v1/dataset5/prepared_180day_source.parquet",
        "protocol_version": "d5_60target_v1",
    },
    "D6": {
        "dataset": "D6",
        "store": "CA_1",
        "group_field": "dept_id",
        "group_value": "FOODS_3",
        "new_root": ROOT / "outputs" / "d6_60target_robustness_20261004",
        "manifest_name": "d6_60target_manifest.json",
        "out_name": "D6_70target_BL_SITA.xlsx",
        "new_manifest_sha256": "6fb0b8045ec912a661fbc2d705ee38ae6429b3d2af3423020b3a09447d30b0f7",
        "source_sha256": "12297d9b00a8e40fd9f13966e41aa558fe15e00079f8d0b23477bf4e7c288458",
        "universe_id": "832c50ec32b11c8f",
        "old_source_data_path": "数据集/固化数据/d1_d6_sealed_v1/dataset6/source.parquet",
        "protocol_version": "d6_60target_v1",
    },
}
NEW_ROOT = DATASETS["D5"]["new_root"]
OUT_XLSX = NEW_ROOT / DATASETS["D5"]["out_name"]
NEW_MANIFEST_SHA = DATASETS["D5"]["new_manifest_sha256"]
SOURCE_SHA = DATASETS["D5"]["source_sha256"]
UNIVERSE = DATASETS["D5"]["universe_id"]
CFG = DATASETS["D5"]


def apply_cfg(dataset: str) -> dict:
    global NEW_ROOT, OUT_XLSX, NEW_MANIFEST_SHA, SOURCE_SHA, UNIVERSE, CFG
    CFG = DATASETS[dataset]
    NEW_ROOT = CFG["new_root"]
    OUT_XLSX = NEW_ROOT / CFG["out_name"]
    NEW_MANIFEST_SHA = CFG["new_manifest_sha256"]
    SOURCE_SHA = CFG["source_sha256"]
    UNIVERSE = CFG["universe_id"]
    return CFG

METHOD_NORM = {
    "SeasonalNaive": "Seasonal Naive",
    "Seasonal Naive": "Seasonal Naive",
    "LightGBM": "LightGBM",
    "LSTM": "LSTM",
    "DLinear": "DLinear",
    "TCN": "TCN",
    "PatchTST": "PatchTST",
    "TimesNet": "TimesNet",
    "MSML-TL": "MSML-TL",
    "SITA": "SITA",
}
METHOD_ORDER = (
    "SITA",
    "Seasonal Naive",
    "LightGBM",
    "LSTM",
    "DLinear",
    "TCN",
    "PatchTST",
    "TimesNet",
    "MSML-TL",
)
OLD_METHODS = {"SITA", "Seasonal Naive", "PatchTST", "TimesNet"}
OLD_FILL_METHODS = ("LightGBM", "LSTM", "DLinear", "TCN", "MSML-TL")
METHOD_SLUGS = {
    "Seasonal Naive": "seasonal_naive",
    "LightGBM": "lightgbm",
    "LSTM": "lstm",
    "DLinear": "dlinear",
    "TCN": "tcn",
    "PatchTST": "patchtst",
    "TimesNet": "timesnet",
    "MSML-TL": "msml_tl",
}
CELL_COLUMNS = [
    "cohort",
    "cohort_label",
    "dataset",
    "store_id",
    "item_id",
    "target_identity",
    "target_rank",
    "selection_origin",
    "method",
    "horizon",
    "seed",
    "n_test",
    "rmse",
    "smape",
    "status",
    "first_target_date",
    "last_target_date",
    "train_start",
    "train_end",
    "val_start",
    "val_end",
    "n_train_days",
    "n_val_days",
    "n_test_days",
    "deterministic_reuse",
    "prediction_source",
    "prediction_finite_count",
    "target_finite_count",
    "selected_sources",
    "source_1",
    "source_1_distance",
    "source_1_weight",
    "source_2",
    "source_2_distance",
    "source_2_weight",
    "source_3",
    "source_3_distance",
    "source_3_weight",
    "cache_status",
    "new_source_training",
    "canary_cell",
    "msml_selected_sources",
    "msml_formal_bl8",
    "msml_used_sita_top3",
    "msml_source_cache_hit",
    "msml_source_cache_miss",
    "msml_used_stage_a_cache",
    "cell_key",
    "key_sha256",
    "manifest_sha256",
    "protocol_sha256",
    "source_universe_id",
    "source_data_sha256",
    "source_data_path",
    "evaluation_dates_sha256",
    "direct_horizon",
    "trained_at",
    "result_dir",
]


def _utc() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _num(value):
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return value
    text = str(value).strip()
    if text.lower() in {"true", "false"}:
        return text.lower() == "true"
    try:
        if "." in text or "e" in text.lower():
            return float(text)
        return int(text)
    except ValueError:
        return value


def _bool(value):
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in {"true", "1"}:
        return True
    if text in {"false", "0"}:
        return False
    return value


def _split_identity(identity: str) -> tuple[str, str]:
    store, item = str(identity).split("/", 1)
    return store, item


def _parse_sources(value) -> list[str]:
    if value in (None, ""):
        return []
    if isinstance(value, list):
        return [str(item) for item in value if item not in (None, "")]
    text = str(value).strip()
    if not text:
        return []
    if text.startswith("["):
        try:
            parsed = ast.literal_eval(text)
            if isinstance(parsed, list):
                return [str(item) for item in parsed if item not in (None, "")]
        except (SyntaxError, ValueError):
            pass
    return [part.strip() for part in text.split(",") if part.strip()]


def _blank_row() -> dict:
    return {key: None for key in CELL_COLUMNS}


def _attach_sources(row: dict, sources: list[str], details: list[dict] | None = None) -> None:
    if sources:
        row["selected_sources"] = ",".join(sources)
    details = details or []
    for index in range(3):
        prefix = f"source_{index + 1}"
        if index < len(details):
            item = details[index]
            row[prefix] = item.get("identity") or item.get("source_identity")
            row[f"{prefix}_distance"] = item.get("distance")
            row[f"{prefix}_weight"] = item.get("weight")
        elif index < len(sources):
            row[prefix] = sources[index]


def _load_new_universe_ids() -> dict[str, str]:
    path = NEW_ROOT / "source_selection_results.csv"
    if not path.is_file():
        return {}
    out = {}
    with path.open(encoding="utf-8", newline="") as handle:
        for raw in csv.DictReader(handle):
            identity = str(raw.get("target_identity") or "")
            universe = str(raw.get("source_universe_id") or "")
            if identity and universe:
                out[identity] = universe
    return out


def _normalize_new_record(item: dict, universe_by_identity: dict[str, str]) -> dict:
    identity = str(item["target_identity"])
    store, _ = _split_identity(identity)
    if item.get("target_rank") not in (None, ""):
        rank = int(item["target_rank"])
    else:
        within = int(item.get("target_rank_within_store") or 0)
        stores = list(CFG.get("stores") or [])
        rank = (stores.index(store) * 10 + within) if store in stores else within
    return {
        "dataset": item.get("dataset") or CFG["dataset"],
        "target_identity": identity,
        "target_rank": rank,
        "target_rank_within_store": item.get("target_rank_within_store"),
        "selection_origin": item.get("selection_origin"),
        "base_sampling_seed": item.get("base_sampling_seed") or item.get("sampling_seed"),
        "eligibility_status": item.get("eligibility_status"),
        "train_start": item.get("train_start") or item.get("target_train_start"),
        "train_end": item.get("train_end") or item.get("target_train_end"),
        "val_start": item.get("val_start") or item.get("target_val_start"),
        "val_end": item.get("val_end") or item.get("target_val_end"),
        "test_start": item.get("test_start") or item.get("target_test_start"),
        "test_end": item.get("test_end") or item.get("target_test_end"),
        "n_train_days": item.get("n_train_days") or 15,
        "n_val_days": item.get("n_val_days") or 15,
        "n_test_days": item.get("n_test_days") or 180,
        "source_universe_id": (
            item.get("source_universe_id")
            or universe_by_identity.get(identity)
            or item.get("source_universe_store")
        ),
        "source_data_path": item.get("source_data_path") or item.get("source_calendar_path") or CFG["old_source_data_path"],
        "source_data_sha256": item.get("source_data_sha256") or item.get("source_calendar_sha256") or SOURCE_SHA,
        "protocol_version": item.get("protocol_version") or CFG.get("protocol_version"),
        "selection_pool_size": item.get("selection_pool_size"),
    }


def load_new_manifest() -> dict[str, dict]:
    records = json.loads((NEW_ROOT / CFG["manifest_name"]).read_text(encoding="utf-8"))["records"]
    universe_by_identity = _load_new_universe_ids()
    return {str(item["target_identity"]): _normalize_new_record(item, universe_by_identity) for item in records}


def load_old_manifest() -> dict[str, dict]:
    records = json.loads((OLD_ROOT / "additional_targets_manifest.json").read_text(encoding="utf-8"))["records"]
    out = {}
    for item in records:
        if item.get("dataset") != CFG["dataset"]:
            continue
        identity = str(item["target_identity"])
        out[identity] = {
            "dataset": CFG["dataset"],
            "target_identity": identity,
            "target_rank": int(item["selected_rank"]),
            "selection_origin": "OLD_ADDITIONAL_10",
            "base_sampling_seed": item.get("target_sampling_seed"),
            "eligibility_status": item.get("eligibility_status"),
            "train_start": item.get("target_train_start"),
            "train_end": item.get("target_train_end"),
            "val_start": item.get("target_val_start"),
            "val_end": item.get("target_val_end"),
            "test_start": item.get("target_test_start"),
            "test_end": item.get("target_test_end"),
            "n_train_days": 15,
            "n_val_days": 15,
            "n_test_days": 180,
            "source_universe_id": str(item.get("source_candidate_pool_id") or UNIVERSE),
            "source_data_path": CFG["old_source_data_path"],
            "source_data_sha256": SOURCE_SHA,
            "protocol_version": "additional_targets_manifest",
            "selection_pool_size": item.get("selection_pool_size"),
        }
    return out


def _fill_manifest(row: dict, meta: dict | None) -> None:
    if not meta:
        return
    row["target_rank"] = meta.get("target_rank")
    row["selection_origin"] = meta.get("selection_origin")
    row["train_start"] = meta.get("train_start")
    row["train_end"] = meta.get("train_end")
    row["val_start"] = meta.get("val_start")
    row["val_end"] = meta.get("val_end")
    row["n_train_days"] = meta.get("n_train_days")
    row["n_val_days"] = meta.get("n_val_days")
    row["n_test_days"] = meta.get("n_test_days")
    row["source_universe_id"] = meta.get("source_universe_id") or row.get("source_universe_id")
    row["source_data_path"] = meta.get("source_data_path")
    row["source_data_sha256"] = meta.get("source_data_sha256") or SOURCE_SHA


def load_new_sita(manifest: dict[str, dict]) -> list[dict]:
    rows = []
    path = NEW_ROOT / "sita" / "sita_progress.csv"
    with path.open(encoding="utf-8", newline="") as handle:
        for raw in csv.DictReader(handle):
            identity = str(raw["target_identity"])
            store, item = _split_identity(identity)
            horizon = int(raw["horizon"])
            seed = int(raw["seed"])
            row = _blank_row()
            row.update(
                {
                    "cohort": "NEW_RANDOM_60",
                    "cohort_label": "new 60-target robustness 20261004",
                    "dataset": CFG["dataset"],
                    "store_id": store,
                    "item_id": item,
                    "target_identity": identity,
                    "method": "SITA",
                    "horizon": horizon,
                    "seed": seed,
                    "n_test": _num(raw.get("n_test")),
                    "rmse": _num(raw.get("rmse")),
                    "smape": _num(raw.get("smape")),
                    "status": raw.get("status"),
                    "first_target_date": raw.get("first_target_date"),
                    "last_target_date": raw.get("last_target_date"),
                    "deterministic_reuse": False,
                    "prediction_source": "trained",
                    "prediction_finite_count": _num(raw.get("prediction_finite_count")),
                    "target_finite_count": _num(raw.get("target_finite_count")),
                    "cache_status": raw.get("cache_status"),
                    "new_source_training": _bool(raw.get("new_source_training")),
                    "canary_cell": _bool(raw.get("canary_cell")),
                    "cell_key": raw.get("cell_key"),
                    "manifest_sha256": raw.get("manifest_sha256") or NEW_MANIFEST_SHA,
                    "protocol_sha256": raw.get("protocol_sha256") or PROTOCOL,
                    "direct_horizon": "y[forecast_origin + horizon]",
                    "trained_at": raw.get("trained_at"),
                    "result_dir": str(
                        (NEW_ROOT / "sita" / "cells" / f"{store}__{item}" / f"h{horizon}" / f"s{seed}").relative_to(ROOT)
                    ),
                }
            )
            _fill_manifest(row, manifest.get(identity))
            sources = _parse_sources(raw.get("selected_sources"))
            usage_path = NEW_ROOT / "sita" / "cells" / f"{store}__{item}" / f"h{horizon}" / f"s{seed}" / "source_usage.json"
            details = []
            if usage_path.is_file():
                details = json.loads(usage_path.read_text(encoding="utf-8"))
            _attach_sources(row, sources, details)
            rows.append(row)
    return rows


def load_new_baselines(manifest: dict[str, dict]) -> list[dict]:
    rows = []
    path = NEW_ROOT / "baselines" / "baseline_progress.csv"
    with path.open(encoding="utf-8", newline="") as handle:
        for raw in csv.DictReader(handle):
            method = METHOD_NORM.get(str(raw.get("method")), str(raw.get("method")))
            identity = str(raw.get("target_identity") or "")
            if not identity:
                continue
            store, item = _split_identity(identity)
            horizon = int(float(raw["horizon"]))
            seed = int(float(raw["seed"]))
            slug = {
                "Seasonal Naive": "seasonal_naive",
                "LightGBM": "lightgbm",
                "LSTM": "lstm",
                "DLinear": "dlinear",
                "TCN": "tcn",
                "PatchTST": "patchtst",
                "TimesNet": "timesnet",
                "MSML-TL": "msml_tl",
            }.get(method, method.lower().replace(" ", "_"))
            row = _blank_row()
            row.update(
                {
                    "cohort": "NEW_RANDOM_60",
                    "cohort_label": "new 60-target robustness 20261004",
                    "dataset": CFG["dataset"],
                    "store_id": store,
                    "item_id": item,
                    "target_identity": identity,
                    "method": method,
                    "horizon": horizon,
                    "seed": seed,
                    "n_test": _num(raw.get("n_test")),
                    "rmse": _num(raw.get("rmse")),
                    "smape": _num(raw.get("smape")),
                    "status": raw.get("status"),
                    "first_target_date": raw.get("first_target_date"),
                    "last_target_date": raw.get("last_target_date"),
                    "deterministic_reuse": _bool(raw.get("deterministic_reuse")),
                    "prediction_source": "computed" if method == "Seasonal Naive" else "trained",
                    "prediction_finite_count": _num(raw.get("prediction_finite_count")),
                    "target_finite_count": _num(raw.get("target_finite_count")),
                    "canary_cell": False,
                    "msml_selected_sources": ",".join(_parse_sources(raw.get("msml_selected_sources"))) or None,
                    "msml_formal_bl8": _bool(raw.get("msml_formal_bl8")),
                    "msml_used_sita_top3": _bool(raw.get("msml_used_sita_top3")),
                    "msml_source_cache_hit": _num(raw.get("msml_source_cache_hit")),
                    "msml_source_cache_miss": _num(raw.get("msml_source_cache_miss")),
                    "msml_used_stage_a_cache": _bool(raw.get("msml_used_stage_a_cache")),
                    "cell_key": raw.get("cell_key") or None,
                    "manifest_sha256": raw.get("manifest_sha256") or NEW_MANIFEST_SHA,
                    "protocol_sha256": raw.get("protocol_sha256") or PROTOCOL,
                    "direct_horizon": "y[forecast_origin + horizon]",
                    "trained_at": raw.get("trained_at"),
                    "result_dir": str(
                        (
                            NEW_ROOT
                            / "baselines"
                            / slug
                            / "cells"
                            / f"{store}_{item}__h{horizon}__s{seed}"
                        ).relative_to(ROOT)
                    ),
                }
            )
            _fill_manifest(row, manifest.get(identity))
            if method == "MSML-TL":
                _attach_sources(row, _parse_sources(raw.get("msml_selected_sources")))
            rows.append(row)
    completed_keys = {
        (row["method"], row["target_identity"], row["horizon"], row["seed"])
        for row in rows
        if str(row.get("status")) in {"COMPLETED", "SKIP"}
    }
    kept = []
    for row in rows:
        key = (row["method"], row["target_identity"], row["horizon"], row["seed"])
        if str(row.get("status")) in {"COMPLETED", "SKIP"}:
            kept.append(row)
        elif key not in completed_keys:
            kept.append(row)
    return kept


def load_old_cells(manifest: dict[str, dict]) -> list[dict]:
    rows = []
    metrics_path = OLD_ROOT / "full_run" / "metrics.csv"
    with metrics_path.open(encoding="utf-8", newline="") as handle:
        for raw in csv.DictReader(handle):
            if raw.get("dataset") != CFG["dataset"]:
                continue
            identity = str(raw["target_identity"])
            method = METHOD_NORM[str(raw["method"])]
            store, item = _split_identity(identity)
            horizon = int(raw["horizon"])
            seed = int(raw["seed"])
            folder = {
                "Seasonal Naive": "SeasonalNaive",
                "PatchTST": "PatchTST",
                "TimesNet": "TimesNet",
                "SITA": "SITA",
            }[method]
            cell_dir = OLD_ROOT / "full_run" / "cells" / CFG["dataset"] / f"{store}__{item}" / f"h{horizon}" / f"s{seed}" / folder
            key_payload = {}
            done_payload = {}
            if (cell_dir / "key.json").is_file():
                key_payload = json.loads((cell_dir / "key.json").read_text(encoding="utf-8"))
            if (cell_dir / "DONE.json").is_file():
                done_payload = json.loads((cell_dir / "DONE.json").read_text(encoding="utf-8"))
            details = list(key_payload.get("selection") or [])
            sources = [str(item_sel.get("identity")) for item_sel in details if item_sel.get("identity")]
            row = _blank_row()
            row.update(
                {
                    "cohort": "OLD_ADDITIONAL_10",
                    "cohort_label": "old additional-target robustness 10 targets",
                    "dataset": CFG["dataset"],
                    "store_id": store,
                    "item_id": item,
                    "target_identity": identity,
                    "method": method,
                    "horizon": horizon,
                    "seed": seed,
                    "n_test": _num(raw.get("n_test")),
                    "rmse": _num(raw.get("rmse")),
                    "smape": _num(raw.get("smape")),
                    "status": done_payload.get("status") or "COMPLETED",
                    "first_target_date": raw.get("first_target_date"),
                    "last_target_date": raw.get("last_target_date"),
                    "deterministic_reuse": _bool(raw.get("deterministic_reuse")),
                    "prediction_source": raw.get("prediction_source"),
                    "prediction_finite_count": _num(raw.get("prediction_finite_count")),
                    "target_finite_count": _num(raw.get("target_finite_count")),
                    "cache_status": None if method != "SITA" else "OFFICIAL_FULL_RUN",
                    "new_source_training": False if method == "SITA" else None,
                    "canary_cell": False,
                    "cell_key": None,
                    "key_sha256": done_payload.get("key_sha256"),
                    "manifest_sha256": key_payload.get("manifest_sha256") or OLD_MANIFEST_SHA,
                    "protocol_sha256": key_payload.get("protocol_sha256") or PROTOCOL,
                    "source_universe_id": raw.get("source_universe_id") or UNIVERSE,
                    "evaluation_dates_sha256": key_payload.get("evaluation_dates_sha256"),
                    "direct_horizon": key_payload.get("direct_horizon") or "y[forecast_origin + horizon]",
                    "result_dir": str(cell_dir.relative_to(ROOT)),
                }
            )
            _fill_manifest(row, manifest.get(identity))
            if method == "SITA":
                _attach_sources(row, sources, details)
            rows.append(row)
    return rows


def load_old_missing_baselines(manifest: dict[str, dict]) -> list[dict]:
    """D4/D5/D6 old-10 fill for LightGBM/LSTM/DLinear/TCN/MSML-TL."""
    if CFG["dataset"] not in {"D4", "D5", "D6"}:
        return []
    path = OLD_FILL_ROOT / CFG["dataset"] / "baselines" / "baseline_progress.csv"
    if not path.is_file():
        return []
    rows = []
    with path.open(encoding="utf-8", newline="") as handle:
        for raw in csv.DictReader(handle):
            method = METHOD_NORM.get(str(raw.get("method")), str(raw.get("method")))
            if method not in OLD_FILL_METHODS:
                continue
            identity = str(raw.get("target_identity") or "")
            if not identity:
                continue
            store, item = _split_identity(identity)
            horizon = int(float(raw["horizon"]))
            seed = int(float(raw["seed"]))
            slug = METHOD_SLUGS[method]
            cell_dir = (
                OLD_FILL_ROOT
                / CFG["dataset"]
                / "baselines"
                / slug
                / "cells"
                / f"{store}_{item}__h{horizon}__s{seed}"
            )
            row = _blank_row()
            row.update(
                {
                    "cohort": "OLD_ADDITIONAL_10",
                    "cohort_label": "old additional-target robustness 10 targets",
                    "dataset": CFG["dataset"],
                    "store_id": store,
                    "item_id": item,
                    "target_identity": identity,
                    "method": method,
                    "horizon": horizon,
                    "seed": seed,
                    "n_test": _num(raw.get("n_test")),
                    "rmse": _num(raw.get("rmse")),
                    "smape": _num(raw.get("smape")),
                    "status": raw.get("status"),
                    "first_target_date": raw.get("first_target_date"),
                    "last_target_date": raw.get("last_target_date"),
                    "deterministic_reuse": _bool(raw.get("deterministic_reuse")),
                    "prediction_source": "computed" if method == "Seasonal Naive" else "trained",
                    "prediction_finite_count": _num(raw.get("prediction_finite_count")),
                    "target_finite_count": _num(raw.get("target_finite_count")),
                    "canary_cell": False,
                    "msml_selected_sources": ",".join(_parse_sources(raw.get("msml_selected_sources"))) or None,
                    "msml_formal_bl8": _bool(raw.get("msml_formal_bl8")),
                    "msml_used_sita_top3": _bool(raw.get("msml_used_sita_top3")),
                    "msml_source_cache_hit": _num(raw.get("msml_source_cache_hit")),
                    "msml_source_cache_miss": _num(raw.get("msml_source_cache_miss")),
                    "msml_used_stage_a_cache": _bool(raw.get("msml_used_stage_a_cache")),
                    "cell_key": raw.get("cell_key") or None,
                    "manifest_sha256": raw.get("manifest_sha256") or OLD_MANIFEST_SHA,
                    "protocol_sha256": raw.get("protocol_sha256") or PROTOCOL,
                    "direct_horizon": "y[forecast_origin + horizon]",
                    "trained_at": raw.get("trained_at"),
                    "result_dir": str(cell_dir.relative_to(ROOT)),
                }
            )
            _fill_manifest(row, manifest.get(identity))
            if method == "MSML-TL":
                _attach_sources(row, _parse_sources(raw.get("msml_selected_sources")))
            rows.append(row)
    completed_keys = {
        (row["method"], row["target_identity"], row["horizon"], row["seed"])
        for row in rows
        if str(row.get("status")) in {"COMPLETED", "SKIP"}
    }
    kept = []
    for row in rows:
        key = (row["method"], row["target_identity"], row["horizon"], row["seed"])
        if str(row.get("status")) in {"COMPLETED", "SKIP"}:
            kept.append(row)
        elif key not in completed_keys:
            kept.append(row)
    return kept


def sort_rows(rows: list[dict]) -> list[dict]:
    method_rank = {name: index for index, name in enumerate(METHOD_ORDER)}
    cohort_rank = {"OLD_ADDITIONAL_10": 0, "NEW_RANDOM_60": 1}

    def key(row: dict):
        return (
            cohort_rank.get(str(row.get("cohort")), 9),
            int(row.get("target_rank") or 10**9),
            str(row.get("target_identity")),
            method_rank.get(str(row.get("method")), 99),
            int(row.get("horizon") or 0),
            int(row.get("seed") or 0),
        )

    return sorted(rows, key=key)


def seed_mean_rows(rows: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for row in rows:
        if str(row.get("status")) not in {"COMPLETED", "SKIP"}:
            continue
        if row.get("rmse") is None:
            continue
        key = (
            row["cohort"],
            row["method"],
            row["target_identity"],
            row["horizon"],
        )
        groups[key].append(row)
    out = []
    for (cohort, method, identity, horizon), items in sorted(groups.items(), key=lambda item: (item[0][0], int(item[0][3]), METHOD_ORDER.index(item[0][1]) if item[0][1] in METHOD_ORDER else 99, item[0][2])):
        rmses = [float(item["rmse"]) for item in items]
        smapes = [float(item["smape"]) for item in items]
        sample = items[0]
        out.append(
            {
                "cohort": cohort,
                "method": method,
                "store_id": sample["store_id"],
                "target_identity": identity,
                "target_rank": sample.get("target_rank"),
                "selection_origin": sample.get("selection_origin"),
                "horizon": horizon,
                "n_seeds": len(items),
                "rmse_mean": statistics.fmean(rmses),
                "rmse_std": statistics.stdev(rmses) if len(rmses) > 1 else 0.0,
                "smape_mean": statistics.fmean(smapes),
                "smape_std": statistics.stdev(smapes) if len(smapes) > 1 else 0.0,
                "n_test": sample.get("n_test"),
                "first_target_date": sample.get("first_target_date"),
                "last_target_date": sample.get("last_target_date"),
                "selected_sources": sample.get("selected_sources"),
                "deterministic_reuse": sample.get("deterministic_reuse"),
            }
        )
    return out


def summary_rows(seed_rows: list[dict]) -> list[dict]:
    groups = defaultdict(list)
    for row in seed_rows:
        groups[(row["cohort"], row["method"], row["horizon"])].append(row)
    old_methods = {row["method"] for row in seed_rows if row["cohort"] == "OLD_ADDITIONAL_10"}
    new_methods = {row["method"] for row in seed_rows if row["cohort"] == "NEW_RANDOM_60"}
    shared = old_methods & new_methods
    both = defaultdict(list)
    for row in seed_rows:
        if row["method"] in shared:
            both[(row["method"], row["horizon"])].append(row)
    out = []
    for (cohort, method, horizon), items in sorted(
        groups.items(),
        key=lambda item: (0 if item[0][0] == "OLD_ADDITIONAL_10" else 1, METHOD_ORDER.index(item[0][1]) if item[0][1] in METHOD_ORDER else 99, item[0][2]),
    ):
        rmses = [float(item["rmse_mean"]) for item in items]
        smapes = [float(item["smape_mean"]) for item in items]
        out.append(
            {
                "cohort": cohort,
                "method": method,
                "horizon": horizon,
                "n_targets": len(items),
                "rmse_mean": statistics.fmean(rmses),
                "rmse_median": statistics.median(rmses),
                "rmse_std": statistics.pstdev(rmses) if len(rmses) > 1 else 0.0,
                "smape_mean": statistics.fmean(smapes),
                "smape_median": statistics.median(smapes),
                "smape_std": statistics.pstdev(smapes) if len(smapes) > 1 else 0.0,
            }
        )
    for (method, horizon), items in sorted(both.items(), key=lambda item: (METHOD_ORDER.index(item[0][0]) if item[0][0] in METHOD_ORDER else 99, item[0][1])):
        rmses = [float(item["rmse_mean"]) for item in items]
        smapes = [float(item["smape_mean"]) for item in items]
        out.append(
            {
                "cohort": "ALL_70",
                "method": method,
                "horizon": horizon,
                "n_targets": len(items),
                "rmse_mean": statistics.fmean(rmses),
                "rmse_median": statistics.median(rmses),
                "rmse_std": statistics.pstdev(rmses) if len(rmses) > 1 else 0.0,
                "smape_mean": statistics.fmean(smapes),
                "smape_median": statistics.median(smapes),
                "smape_std": statistics.pstdev(smapes) if len(smapes) > 1 else 0.0,
            }
        )
    return out


def coverage_rows(rows: list[dict]) -> list[dict]:
    counts = defaultdict(int)
    present = defaultdict(set)
    for row in rows:
        counts[(row["cohort"], row["method"], str(row.get("status")))] += 1
        present[row["cohort"]].add(row["method"])
    out = []
    for cohort in ("OLD_ADDITIONAL_10", "NEW_RANDOM_60"):
        expected_targets = 10 if cohort == "OLD_ADDITIONAL_10" else 60
        if cohort == "NEW_RANDOM_60":
            methods = set(METHOD_ORDER)
        else:
            methods = set(OLD_METHODS) | (present[cohort] & set(OLD_FILL_METHODS))
        for method in METHOD_ORDER:
            if method not in methods:
                continue
            completed = counts[(cohort, method, "COMPLETED")] + counts[(cohort, method, "SKIP")]
            expected = expected_targets * 3 * 5
            failed = sum(
                n for (c, m, status), n in counts.items() if c == cohort and m == method and status not in {"COMPLETED", "SKIP"}
            )
            out.append(
                {
                    "cohort": cohort,
                    "method": method,
                    "n_targets_expected": expected_targets,
                    "n_cells_expected": expected,
                    "n_cells_completed": completed,
                    "n_cells_other": failed,
                    "complete": completed == expected and failed == 0,
                }
            )
    return out


def target_rows(old_meta: dict[str, dict], new_meta: dict[str, dict]) -> list[dict]:
    out = []
    for meta in sorted(old_meta.values(), key=lambda item: int(item["target_rank"])):
        row = dict(meta)
        row["cohort"] = "OLD_ADDITIONAL_10"
        out.append(row)
    for meta in sorted(new_meta.values(), key=lambda item: int(item["target_rank"])):
        row = {
            "cohort": "NEW_RANDOM_60",
            "dataset": meta.get("dataset"),
            "target_identity": meta.get("target_identity"),
            "target_rank": meta.get("target_rank"),
            "selection_origin": meta.get("selection_origin"),
            "base_sampling_seed": meta.get("base_sampling_seed"),
            "eligibility_status": meta.get("eligibility_status"),
            "train_start": meta.get("train_start"),
            "train_end": meta.get("train_end"),
            "val_start": meta.get("val_start"),
            "val_end": meta.get("val_end"),
            "test_start": meta.get("test_start"),
            "test_end": meta.get("test_end"),
            "n_train_days": meta.get("n_train_days"),
            "n_val_days": meta.get("n_val_days"),
            "n_test_days": meta.get("n_test_days"),
            "source_universe_id": meta.get("source_universe_id"),
            "source_data_path": meta.get("source_data_path"),
            "source_data_sha256": meta.get("source_data_sha256"),
            "protocol_version": meta.get("protocol_version"),
        }
        out.append(row)
    return out


def _write_sheet(wb: Workbook, name: str, rows: list[dict], columns: list[str]) -> None:
    ws = wb.create_sheet(name)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    ws.append(columns)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append([row.get(column) for column in columns])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(1, len(rows) + 1)}"
    ws.freeze_panes = "A2"
    for index, column in enumerate(columns, start=1):
        width = min(max(len(column) + 2, 12), 42)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.auto_filter.ref = ws.auto_filter.ref


def build(dataset: str = "D5") -> Path:
    apply_cfg(dataset)
    new_meta = load_new_manifest()
    old_meta = load_old_manifest()
    all_rows = sort_rows(
        load_old_cells(old_meta)
        + load_old_missing_baselines(old_meta)
        + load_new_sita(new_meta)
        + load_new_baselines(new_meta)
    )
    failures = [row for row in all_rows if str(row.get("status")) not in {"COMPLETED", "SKIP"}]
    rows = [row for row in all_rows if str(row.get("status")) in {"COMPLETED", "SKIP"}]
    seed_rows = seed_mean_rows(rows)
    summaries = summary_rows(seed_rows)
    coverage = coverage_rows(rows)
    targets = target_rows(old_meta, new_meta)
    n_old = sum(1 for row in rows if row["cohort"] == "OLD_ADDITIONAL_10")
    n_new = sum(1 for row in rows if row["cohort"] == "NEW_RANDOM_60")
    n_sita_new = sum(1 for row in rows if row["cohort"] == "NEW_RANDOM_60" and row["method"] == "SITA")
    n_bl_new = n_new - n_sita_new
    n_old_fill = sum(1 for row in rows if row["cohort"] == "OLD_ADDITIONAL_10" and row["method"] in OLD_FILL_METHODS)
    old_methods_present = sorted(
        {row["method"] for row in rows if row["cohort"] == "OLD_ADDITIONAL_10"},
        key=lambda name: METHOD_ORDER.index(name) if name in METHOD_ORDER else 99,
    )

    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "cell_raw", rows, CELL_COLUMNS)
    _write_sheet(
        wb,
        "seed_mean",
        seed_rows,
        [
            "cohort",
            "method",
            "store_id",
            "target_identity",
            "target_rank",
            "selection_origin",
            "horizon",
            "n_seeds",
            "rmse_mean",
            "rmse_std",
            "smape_mean",
            "smape_std",
            "n_test",
            "first_target_date",
            "last_target_date",
            "selected_sources",
            "deterministic_reuse",
        ],
    )
    _write_sheet(
        wb,
        "method_horizon_summary",
        summaries,
        [
            "cohort",
            "method",
            "horizon",
            "n_targets",
            "rmse_mean",
            "rmse_median",
            "rmse_std",
            "smape_mean",
            "smape_median",
            "smape_std",
        ],
    )
    _write_sheet(
        wb,
        "coverage",
        coverage,
        [
            "cohort",
            "method",
            "n_targets_expected",
            "n_cells_expected",
            "n_cells_completed",
            "n_cells_other",
            "complete",
        ],
    )
    target_columns = [
        "cohort",
        "dataset",
        "target_identity",
        "target_rank",
        "selection_origin",
        "base_sampling_seed",
        "eligibility_status",
        "train_start",
        "train_end",
        "val_start",
        "val_end",
        "test_start",
        "test_end",
        "n_train_days",
        "n_val_days",
        "n_test_days",
        "source_universe_id",
        "source_data_path",
        "source_data_sha256",
        "protocol_version",
        "selection_pool_size",
    ]
    _write_sheet(wb, "targets", targets, target_columns)
    if failures:
        _write_sheet(wb, "failures", failures, CELL_COLUMNS)
    readme = [
        {"field": "dataset", "value": CFG["dataset"]},
        {"field": "store", "value": CFG["store"]},
        {"field": CFG["group_field"], "value": CFG["group_value"]},
        {"field": "experiment", "value": f"{CFG['dataset']} new 60-target BL+SITA integrated with old 10 additional-target BL+SITA"},
        {"field": "n_targets_total", "value": 70},
        {"field": "n_targets_old_10", "value": 10},
        {"field": "n_targets_new_60", "value": 60},
        {"field": "horizons", "value": "1, 3, 5"},
        {"field": "seeds", "value": "42, 43, 44, 45, 46"},
        {"field": "n_test", "value": 180},
        {"field": "evaluation_index", "value": "direct-h; label y[forecast_origin + horizon]"},
        {"field": "n_cells_old_10", "value": n_old},
        {"field": "n_cells_old_10_fill_5bl", "value": n_old_fill},
        {"field": "n_cells_new_60", "value": n_new},
        {"field": "n_cells_new_sita", "value": n_sita_new},
        {"field": "n_cells_new_baselines", "value": n_bl_new},
        {"field": "n_cells_total", "value": len(rows)},
        {"field": "n_failure_rows_superseded", "value": len(failures)},
        {
            "field": "old_10_methods",
            "value": ", ".join(old_methods_present),
        },
        {
            "field": "new_60_methods",
            "value": "SITA + Seasonal Naive, LightGBM, LSTM, DLinear, TCN, PatchTST, TimesNet, MSML-TL",
        },
        {
            "field": "old_10_note",
            "value": (
                "official full_run had SITA/Seasonal Naive/PatchTST/TimesNet; "
                "LightGBM/LSTM/DLinear/TCN/MSML-TL filled from old10_missing_baselines"
                if n_old_fill
                else "old 10-target official run did not include LightGBM/LSTM/DLinear/TCN/MSML-TL"
            ),
        },
        {"field": "old_manifest_sha256", "value": OLD_MANIFEST_SHA},
        {"field": "new_manifest_sha256", "value": NEW_MANIFEST_SHA},
        {"field": "protocol_sha256", "value": PROTOCOL},
        {"field": "source_data_sha256", "value": SOURCE_SHA},
        {"field": "source_universe_id", "value": UNIVERSE},
        {"field": "old_source_universe_id", "value": CFG.get("old_source_universe_id") or ""},
        {"field": "identity_overlap_note", "value": CFG.get("identity_overlap_note") or "old 10 and new 60 are disjoint"},
        {"field": "old_source_dir", "value": "outputs/additional_target_robustness_20261004/full_run"},
        {
            "field": "old_fill_dir",
            "value": (
                str((OLD_FILL_ROOT / CFG["dataset"] / "baselines").relative_to(ROOT))
                if n_old_fill
                else ""
            ),
        },
        {"field": "new_sita_dir", "value": str((NEW_ROOT / "sita").relative_to(ROOT))},
        {"field": "new_baseline_dir", "value": str((NEW_ROOT / "baselines").relative_to(ROOT))},
        {"field": "sheet_cell_raw", "value": "one row per method x target x horizon x seed; old 10 then new 60"},
        {"field": "sheet_seed_mean", "value": "mean/std over seeds 42-46 for each cohort x method x target x horizon"},
        {
            "field": "sheet_method_horizon_summary",
            "value": "cohort-wise summaries plus ALL_70 for methods shared by old 10 and new 60",
        },
        {"field": "written_at", "value": _utc()},
    ]
    _write_sheet(wb, "readme", readme, ["field", "value"])
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    return OUT_XLSX


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset", choices=("D4", "D5", "D6"), default="D5")
    args = parser.parse_args()
    path = build(args.dataset)
    print(path)
