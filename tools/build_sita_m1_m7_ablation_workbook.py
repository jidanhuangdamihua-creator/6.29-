#!/usr/bin/env python3
"""Aggregate SITA M1–M7 ablation condition outputs into one Excel workbook."""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "configs" / "sita_m1_m7_ablation_manifest.json"
CONDITIONS_DIR = ROOT / "outputs" / "sita_m1_m7_ablation" / "conditions"
OUTPUT = ROOT / "outputs" / "sita_m1_m7_ablation" / "sita_m1_m7_ablation_summary.xlsx"

GROUPS = ("M1", "M2", "M3", "M4", "M5", "M6", "M7")
GROUP_SPECS = {
    "M1": {"init_mode": "native", "conv_frozen": False, "revin_enabled": False},
    "M2": {"init_mode": "fused", "conv_frozen": True, "revin_enabled": False},
    "M3": {"init_mode": "fused", "conv_frozen": False, "revin_enabled": False},
    "M4": {"init_mode": "native", "conv_frozen": False, "revin_enabled": True},
    "M5": {"init_mode": "fused", "conv_frozen": True, "revin_enabled": True},
    "M6": {"init_mode": "fused", "conv_frozen": False, "revin_enabled": True},
    "M7": {"init_mode": "nearest", "conv_frozen": False, "revin_enabled": True},
}
CKA_LAYERS = ("conv1", "conv2", "conv3", "flatten")
PAIR_CONTRASTS = (
    ("M7_vs_M1", "M7", "M1"),
    ("M6_vs_M5", "M6", "M5"),
    ("M3_vs_M2", "M3", "M2"),
    ("M4_vs_M1", "M4", "M1"),
)

header_fill = PatternFill("solid", fgColor="1F4E78")
positive_fill = PatternFill("solid", fgColor="E2F0D9")
negative_fill = PatternFill("solid", fgColor="FCE4D6")
thin_gray = Side(style="thin", color="D9E1F2")
border = Border(bottom=thin_gray)
header_font = Font(name="Arial", bold=True, color="FFFFFF")
body_font = Font(name="Arial", size=10, color="000000")
title_font = Font(name="Arial", bold=True, size=14, color="1F1F1F")


def read_json(path: Path) -> Dict[str, Any]:
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def col_letter(index: int) -> str:
    out = ""
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def parse_condition_dir(path: Path) -> Dict[str, Any]:
    parts = path.name.split("__")
    if len(parts) != 4 or not parts[2].startswith("h") or not parts[3].startswith("s"):
        raise RuntimeError(f"Unexpected condition directory name: {path.name}")
    return {
        "dataset": parts[0],
        "target": parts[1].replace("-", "/"),
        "horizon": int(parts[2][1:]),
        "seed": int(parts[3][1:]),
        "condition_dir": path.name,
        "condition_path": str(path),
    }


def add_table(ws, headers: Sequence[str], rows: Sequence[Mapping[str, Any]], table_name: str) -> None:
    ws.append(list(headers))
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(vertical="top")
    ref = f"A1:{col_letter(len(headers))}{ws.max_row}"
    tab = Table(displayName=table_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tab)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ref
    for col_cells in ws.columns:
        letter = col_cells[0].column_letter
        max_len = max(len(str(cell.value or "")) for cell in col_cells[:80])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 42)
    ws.row_dimensions[1].height = 34


def mean(values: Iterable[float]) -> float:
    items = list(values)
    return sum(items) / len(items) if items else float("nan")


def load_rows() -> tuple[List[Dict[str, Any]], List[Dict[str, Any]], List[Dict[str, Any]], List[Dict[str, Any]]]:
    with MANIFEST.open(encoding="utf-8") as handle:
        manifest = json.load(handle)

    condition_dirs = sorted(CONDITIONS_DIR.iterdir())
    if not condition_dirs:
        raise RuntimeError(f"No condition directories found under {CONDITIONS_DIR}")
    if len(condition_dirs) != 60:
        raise RuntimeError(f"Expected 60 condition directories, found {len(condition_dirs)}")

    condition_rows: List[Dict[str, Any]] = []
    detail_rows: List[Dict[str, Any]] = []
    cka_rows: List[Dict[str, Any]] = []

    for condition_dir in condition_dirs:
        if not condition_dir.is_dir():
            continue
        base = parse_condition_dir(condition_dir)
        contract = read_json(condition_dir / "condition_contract.json")
        status = read_json(condition_dir / "condition_status.json")
        validation = read_json(condition_dir / "validation_report.json")
        cka = read_json(condition_dir / "cka.json") if (condition_dir / "cka.json").is_file() else {}

        condition_rows.append(
            {
                **base,
                "scenario": manifest.get("scenario", ""),
                "target_learning_rate": contract.get("target_learning_rate", ""),
                "source_learning_rate": contract.get("source_learning_rate", ""),
                "condition_status": status.get("status", ""),
                "validation_status": validation.get("status", ""),
                "elapsed_seconds": status.get("elapsed_seconds", ""),
                "groups_complete": ",".join(validation.get("groups_complete", [])),
            }
        )

        for group in GROUPS:
            group_dir = condition_dir / "groups" / group
            metrics = read_json(group_dir / "metrics.json")
            group_status = read_json(group_dir / "status.json")
            spec = GROUP_SPECS[group]
            detail_rows.append(
                {
                    **base,
                    "scenario": manifest.get("scenario", ""),
                    "target_learning_rate": contract.get("target_learning_rate", ""),
                    "group": group,
                    "init_mode": spec["init_mode"],
                    "conv_frozen": spec["conv_frozen"],
                    "revin_enabled": spec["revin_enabled"],
                    "group_status": group_status.get("status", ""),
                    "selected_epoch": metrics.get("selected_epoch", ""),
                    "selection_rule": metrics.get("selection_rule", ""),
                    "weight_variant": metrics.get("weight_variant", ""),
                    "metric_space": metrics.get("metric_space", ""),
                    "rmse_original": metrics.get("rmse_original", ""),
                    "smape_original": metrics.get("smape_original", ""),
                    "rmse_normalized": metrics.get("rmse_normalized", ""),
                    "smape_normalized": metrics.get("smape_normalized", ""),
                    "elapsed_seconds": group_status.get("elapsed_seconds", ""),
                    "optimizer_update_count": group_status.get("optimizer_update_count", ""),
                    "ema_update_count": group_status.get("ema_update_count", ""),
                    "reload_max_abs": validation.get("reload", {}).get(group, {}).get("max_abs", ""),
                    "metric_recompute_rmse_abs": validation.get("metric_recompute", {}).get(group, {}).get("rmse_abs", ""),
                    "metric_recompute_smape_abs": validation.get("metric_recompute", {}).get(group, {}).get("smape_abs", ""),
                }
            )

        for comparison in cka.get("comparisons", []):
            row = {
                **base,
                "comparison": comparison.get("comparison", ""),
                "left_group": comparison.get("left", {}).get("group", ""),
                "left_state": comparison.get("left", {}).get("state", ""),
                "right_group": comparison.get("right", {}).get("group", ""),
                "right_state": comparison.get("right", {}).get("state", ""),
                "notes": "; ".join(comparison.get("notes", [])),
            }
            for layer in CKA_LAYERS:
                layer_payload = comparison.get("layers", {}).get(layer, {})
                row[f"cka_{layer}"] = layer_payload.get("value", "")
                row[f"cka_{layer}_status"] = layer_payload.get("status", "")
            cka_rows.append(row)

    detail_rows.sort(key=lambda row: (row["dataset"], row["target"], row["horizon"], row["seed"], row["group"]))
    condition_rows.sort(key=lambda row: (row["dataset"], row["target"], row["horizon"], row["seed"]))
    cka_rows.sort(key=lambda row: (row["dataset"], row["target"], row["horizon"], row["seed"], row["comparison"]))

    summary_rows: List[Dict[str, Any]] = []
    grouped: Dict[tuple[str, str, int], List[Dict[str, Any]]] = defaultdict(list)
    for row in detail_rows:
        grouped[(row["dataset"], row["target"], row["horizon"])].append(row)

    for (dataset, target, horizon), rows in sorted(grouped.items()):
        summary = {
            "dataset": dataset,
            "target": target,
            "horizon": horizon,
            "n_seeds": len({row["seed"] for row in rows}),
        }
        for group in GROUPS:
            group_rows = [row for row in rows if row["group"] == group]
            summary[f"{group}_RMSE_mean"] = mean(float(row["rmse_original"]) for row in group_rows)
            summary[f"{group}_sMAPE_mean"] = mean(float(row["smape_original"]) for row in group_rows)
        summary_rows.append(summary)

    return condition_rows, detail_rows, summary_rows, cka_rows


def pair_headers(left: str, right: str) -> List[str]:
    return [
        "dataset", "target", "horizon", "seed", "left_group", "right_group",
        f"{left}_RMSE_original", f"{right}_RMSE_original",
        "left_minus_right_RMSE", "left_vs_right_RMSE",
        f"{left}_sMAPE_original", f"{right}_sMAPE_original",
        "left_minus_right_sMAPE", "left_vs_right_sMAPE",
    ]


def build_pair_rows(detail_rows: Sequence[Mapping[str, Any]], left: str, right: str) -> List[Dict[str, Any]]:
    index: Dict[tuple[str, str, int, int], Dict[str, Mapping[str, Any]]] = defaultdict(dict)
    for row in detail_rows:
        key = (row["dataset"], row["target"], row["horizon"], row["seed"])
        index[key][row["group"]] = row

    pair_rows: List[Dict[str, Any]] = []
    for key in sorted(index):
        dataset, target, horizon, seed = key
        left_row = index[key].get(left)
        right_row = index[key].get(right)
        if not left_row or not right_row:
            continue
        left_rmse = float(left_row["rmse_original"])
        right_rmse = float(right_row["rmse_original"])
        left_smape = float(left_row["smape_original"])
        right_smape = float(right_row["smape_original"])
        rmse_delta = left_rmse - right_rmse
        smape_delta = left_smape - right_smape
        pair_rows.append(
            {
                "dataset": dataset,
                "target": target,
                "horizon": horizon,
                "seed": seed,
                "left_group": left,
                "right_group": right,
                f"{left}_RMSE_original": left_rmse,
                f"{right}_RMSE_original": right_rmse,
                "left_minus_right_RMSE": rmse_delta,
                "left_vs_right_RMSE": "胜" if rmse_delta < 0 else "负" if rmse_delta > 0 else "平",
                f"{left}_sMAPE_original": left_smape,
                f"{right}_sMAPE_original": right_smape,
                "left_minus_right_sMAPE": smape_delta,
                "left_vs_right_sMAPE": "胜" if smape_delta < 0 else "负" if smape_delta > 0 else "平",
            }
        )
    return pair_rows


def style_numeric_columns(ws, columns: Sequence[int]) -> None:
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in columns and isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000"


def main() -> None:
    condition_rows, detail_rows, summary_rows, cka_rows = load_rows()

    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "说明"
    ws_condition = wb.create_sheet("条件清单")
    ws_detail = wb.create_sheet("M1-M7逐条件逐Seed")
    ws_summary = wb.create_sheet("条件均值汇总")
    ws_m7_m1 = wb.create_sheet("M7-vs-M1逐Seed")
    ws_m6_m5 = wb.create_sheet("M6-vs-M5逐Seed")
    ws_cka = wb.create_sheet("CKA逐条件")

    ws_readme["A1"] = "SITA M1–M7 Ablation 汇总"
    ws_readme["A1"].font = title_font
    ws_readme.merge_cells("A1:D1")
    readme_rows = [
        ("用途", "将 60 个 SITA M1–M7 条件输出聚合为一个 Excel，便于检查逐条件逐 seed 指标、条件均值与关键配对对比。"),
        ("源目录", str(CONDITIONS_DIR)),
        ("源 manifest", str(MANIFEST)),
        ("输出文件", str(OUTPUT)),
        ("范围", "D1/D4/D5/D6；H1/H3/H5；seeds 42–46；groups M1–M7。"),
        ("条件数量", len(condition_rows)),
        ("明细数量", f"{len(detail_rows)} = 60 conditions × 7 groups"),
        ("RMSE 差值定义", "left_minus_right_RMSE = left_RMSE_original − right_RMSE_original；负值表示 left 组更优。"),
        ("状态说明", "condition_status=ALL_REQUIRED_ARTIFACTS_COMPLETE 或 COMPLETE 均表示产物齐全；CLI 仅对 COMPLETE 返回 0。"),
    ]
    ws_readme.append(["项目", "内容"])
    for row in readme_rows:
        ws_readme.append(row)
    for cell in ws_readme[2]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row in ws_readme.iter_rows(min_row=3):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
    ws_readme.column_dimensions["A"].width = 22
    ws_readme.column_dimensions["B"].width = 110
    ws_readme.freeze_panes = "A3"

    condition_headers = [
        "dataset", "target", "horizon", "seed", "scenario", "target_learning_rate",
        "source_learning_rate", "condition_status", "validation_status", "elapsed_seconds",
        "groups_complete", "condition_dir", "condition_path",
    ]
    detail_headers = [
        "dataset", "target", "horizon", "seed", "scenario", "target_learning_rate", "group",
        "init_mode", "conv_frozen", "revin_enabled", "group_status", "selected_epoch",
        "selection_rule", "weight_variant", "metric_space", "rmse_original", "smape_original",
        "rmse_normalized", "smape_normalized", "elapsed_seconds", "optimizer_update_count",
        "ema_update_count", "reload_max_abs", "metric_recompute_rmse_abs",
        "metric_recompute_smape_abs", "condition_dir", "condition_path",
    ]
    summary_headers = ["dataset", "target", "horizon", "n_seeds"] + [
        f"{group}_{metric}" for group in GROUPS for metric in ("RMSE_mean", "sMAPE_mean")
    ]
    cka_headers = [
        "dataset", "target", "horizon", "seed", "comparison", "left_group", "left_state",
        "right_group", "right_state", "notes",
    ] + [f"cka_{layer}" for layer in CKA_LAYERS] + [f"cka_{layer}_status" for layer in CKA_LAYERS]

    add_table(ws_condition, condition_headers, condition_rows, "ConditionList")
    add_table(ws_detail, detail_headers, detail_rows, "M1M7Detail")
    add_table(ws_summary, summary_headers, summary_rows, "ConditionSummary")

    m7_m1_rows = build_pair_rows(detail_rows, "M7", "M1")
    m6_m5_rows = build_pair_rows(detail_rows, "M6", "M5")
    add_table(ws_m7_m1, pair_headers("M7", "M1"), m7_m1_rows, "M7M1Pair")
    add_table(ws_m6_m5, pair_headers("M6", "M5"), m6_m5_rows, "M6M5Pair")
    add_table(ws_cka, cka_headers, cka_rows, "CKAByCondition")

    style_numeric_columns(ws_detail, [16, 17, 18, 19, 20, 23, 24, 25])
    style_numeric_columns(ws_summary, list(range(5, 5 + len(GROUPS) * 2)))
    style_numeric_columns(ws_m7_m1, [7, 8, 9, 11, 12, 13])
    style_numeric_columns(ws_m6_m5, [7, 8, 9, 11, 12, 13])
    style_numeric_columns(ws_cka, list(range(10, 10 + len(CKA_LAYERS))))

    for ws in (ws_m7_m1, ws_m6_m5):
        ws.conditional_formatting.add(f"I2:I{ws.max_row}", CellIsRule(operator="lessThan", formula=["0"], fill=positive_fill))
        ws.conditional_formatting.add(f"I2:I{ws.max_row}", CellIsRule(operator="greaterThan", formula=["0"], fill=negative_fill))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.properties.creator = "Codex"
    wb.properties.title = "SITA M1-M7 ablation summary"
    wb.properties.subject = "Read-only aggregation of 60 formal M1-M7 condition directories"
    wb.save(OUTPUT)

    print(
        json.dumps(
            {
                "output": str(OUTPUT),
                "conditions": len(condition_rows),
                "detail_rows": len(detail_rows),
                "summary_rows": len(summary_rows),
                "m7_m1_pairs": len(m7_m1_rows),
                "m6_m5_pairs": len(m6_m5_rows),
                "cka_rows": len(cka_rows),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
