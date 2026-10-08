#!/usr/bin/env python3
"""Build a non-destructive integrated workbook for the BL8 comparison.

The paper freeze workbook is read as an authority input and never overwritten.
The output is a new copy with a method column added to the SITA raw sheet and
one normalized 120-row table containing SITA and BL8 cell-level records.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path(
    "/Users/ming/Desktop/复现实验/论文文稿相关/"
    "MS_KATL_Paper_Data_Freeze_v1_4_STATS_FROZEN_BL6_BL7_INTEGRATED.xlsx"
)
DEFAULT_BL8_ROOT = ROOT / "results" / "formal_bl8_msml_tl"
DEFAULT_OUTPUT = ROOT / "outputs" / "formal_bl8_msml_tl" / "BL8_MSML_TL_FORMAL_DATA_INTEGRATED.xlsx"
SITA_SHEET = "03_V3_Raw_5Seed"
EXPECTED_DATASETS = {"D1", "D4", "D5", "D6"}
EXPECTED_TARGETS = {
    "D1": "1/10",
    "D4": "166/258",
    "D5": "48/1159414",
    "D6": "CA_1/FOODS_3_080",
}
EXPECTED_HORIZONS = {1, 3, 5}
EXPECTED_SEEDS = {42, 43, 44, 45, 46}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def expected_keys() -> set[tuple[str, str, int, int]]:
    return {
        (dataset, EXPECTED_TARGETS[dataset], horizon, seed)
        for dataset in EXPECTED_DATASETS
        for horizon in EXPECTED_HORIZONS
        for seed in EXPECTED_SEEDS
    }


def validate_matrix(frame: pd.DataFrame, label: str) -> None:
    required = {"dataset", "target", "horizon", "seed", "rmse", "smape"}
    missing = required - set(frame.columns)
    if missing:
        raise ValueError(f"{label} missing normalized columns: {sorted(missing)}")
    keys = list(zip(frame["dataset"], frame["target"], frame["horizon"], frame["seed"]))
    key_set = set(keys)
    if len(frame) != 60 or len(key_set) != 60 or key_set != expected_keys():
        raise ValueError(
            f"{label} is not the exact 60-cell matrix: rows={len(frame)} "
            f"unique={len(key_set)} missing={len(expected_keys() - key_set)} "
            f"unexpected={len(key_set - expected_keys())}"
        )


def load_sita(source: Path) -> tuple[pd.DataFrame, list[str]]:
    frame = pd.read_excel(source, sheet_name=SITA_SHEET)
    frame.columns = [str(column).strip() for column in frame.columns]
    if "method" not in frame.columns:
        frame.insert(2, "method", "SITA")
    else:
        frame["method"] = frame["method"].astype(str).replace({"MS-KATL-V3": "SITA"})
    if "rmse_original" not in frame.columns or "smape_original" not in frame.columns:
        raise ValueError("SITA sheet lacks rmse_original/smape_original")
    normalized = pd.DataFrame(
        {
            "method": "SITA",
            "method_id": "SITA",
            "dataset": frame["dataset"].astype(str),
            "target": frame["target"].astype(str),
            "mode": frame["mode"].astype(str),
            "horizon": frame["horizon"].astype(int),
            "seed": frame["seed"].astype(int),
            "rmse": frame["rmse_original"].astype(float),
            "smape": frame["smape_original"].astype(float),
            "n": pd.NA,
            "metric_space": frame["metric_space"].astype(str),
            "status": frame["status"].astype(str),
            "authority": "paper_freeze",
            "source_file": frame["source_file"].astype(str),
            "source_record": SITA_SHEET,
        }
    )
    validate_matrix(normalized, "SITA")
    return normalized, list(frame.columns)


def load_bl8(bl8_root: Path) -> pd.DataFrame:
    records = []
    for path in sorted((bl8_root / "cells").glob("*/cell_manifest.json")):
        payload = json.loads(path.read_text(encoding="utf-8"))
        if str(payload.get("cell_status", "")).upper() != "PASS":
            raise ValueError(f"BL8 cell is not PASS: {path}")
        records.append(
            {
                "method": str(payload["method"]),
                "method_id": str(payload["method_id"]),
                "dataset": str(payload["dataset"]),
                "target": str(payload["target"]),
                "mode": str(payload["mode"]),
                "horizon": int(payload["horizon"]),
                "seed": int(payload["seed"]),
                "rmse": float(payload["rmse"]),
                "smape": float(payload["smape"]),
                "n": int(payload["n"]),
                "metric_space": str(payload["metric_space"]),
                "status": str(payload["cell_status"]),
                "authority": "formal_bl8",
                "source_file": str(path),
                "source_record": "cell_manifest.json",
            }
        )
    frame = pd.DataFrame.from_records(records)
    validate_matrix(frame, "BL8")
    if not (frame["method"] == "MSML-TL").all():
        raise ValueError("BL8 matrix contains an unexpected method")
    return frame


def set_table_style(table: Table) -> None:
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )


def write_frame_sheet(ws, frame: pd.DataFrame, table_name: str) -> None:
    headers = list(frame.columns)
    ws.append(headers)
    for value in headers:
        cell = ws.cell(row=1, column=headers.index(value) + 1)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in frame.itertuples(index=False, name=None):
        ws.append([None if pd.isna(value) else value for value in row])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center")
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=table_name, ref=ref)
    set_table_style(table)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ref
    for column_cells in ws.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 10), 48)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def add_method_to_sita_sheet(ws) -> None:
    headers = [str(ws.cell(row=1, column=i).value or "").strip() for i in range(1, ws.max_column + 1)]
    dataset_col = headers.index("dataset") + 1
    data_rows = [
        row
        for row in range(2, ws.max_row + 1)
        if ws.cell(row=row, column=dataset_col).value not in (None, "")
    ]
    if not data_rows:
        raise ValueError(f"{SITA_SHEET} has no data rows")
    last_data_row = max(data_rows)
    if "method" not in headers:
        insert_at = 3
        ws.insert_cols(insert_at, 1)
        ws.cell(row=1, column=insert_at, value="method")
        for row in range(2, last_data_row + 1):
            ws.cell(row=row, column=insert_at, value="SITA")
        source_header = ws.cell(row=1, column=insert_at + 1)
        target_header = ws.cell(row=1, column=insert_at)
        target_header._style = copy.copy(source_header._style)
        target_header.number_format = source_header.number_format
        for row in range(2, last_data_row + 1):
            ws.cell(row=row, column=insert_at)._style = copy.copy(ws.cell(row=row, column=insert_at + 1)._style)
        ws.cell(row=1, column=insert_at).font = copy.copy(source_header.font)
    else:
        method_col = headers.index("method") + 1
        for row in range(2, last_data_row + 1):
            ws.cell(row=row, column=method_col, value="SITA")
    ref = f"A1:{get_column_letter(ws.max_column)}{last_data_row}"
    for table in ws.tables.values():
        table.ref = ref


def build(source: Path, bl8_root: Path, output: Path) -> dict[str, object]:
    if not source.exists():
        raise FileNotFoundError(source)
    if not (bl8_root / "cells").exists():
        raise FileNotFoundError(bl8_root / "cells")

    source_hash_before = sha256_file(source)
    sita, _ = load_sita(source)
    bl8 = load_bl8(bl8_root)
    unified = pd.concat([sita, bl8], ignore_index=True)
    key_cols = ["method", "dataset", "target", "mode", "horizon", "seed"]
    if unified.duplicated(key_cols).any():
        raise ValueError("Unified table contains duplicate method/cell keys")
    if len(unified) != 120:
        raise ValueError(f"Expected 120 unified rows, got {len(unified)}")

    output.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(source)
    add_method_to_sita_sheet(wb[SITA_SHEET])
    for sheet_name in ("Integration_Readme", "Unified_Results", "BL8_Raw_60Cell"):
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]

    readme = wb.create_sheet("Integration_Readme", 0)
    readme_rows = [
        ("artifact", "BL8_MSML_TL_FORMAL_DATA_INTEGRATED.xlsx"),
        ("purpose", "SITA and BL8 formal cell-level data in one normalized table"),
        ("source_freeze_workbook", str(source)),
        ("source_freeze_sha256_before", source_hash_before),
        ("source_freeze_sha256_after_read", sha256_file(source)),
        ("source_sita_sheet", SITA_SHEET),
        ("sita_rows", 60),
        ("bl8_rows", 60),
        ("unified_rows", 120),
        ("unified_key", "method + dataset + target + mode + horizon + seed"),
        ("sita_method_repair", "Added method=SITA to the copied 03_V3_Raw_5Seed sheet"),
        ("bl8_source_root", str(bl8_root)),
        ("data_policy", "No averaging or conflict resolution; metric space and authority retained"),
    ]
    readme.append(["field", "value"])
    for row in readme_rows:
        readme.append(list(row))
    for row in readme.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    readme.freeze_panes = "A2"
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["B"].width = 110

    write_frame_sheet(wb.create_sheet("Unified_Results"), unified, "UnifiedResultsTable")
    write_frame_sheet(wb.create_sheet("BL8_Raw_60Cell"), bl8, "BL8Raw60CellTable")

    wb.save(output)
    source_hash_after = sha256_file(source)
    if source_hash_before != source_hash_after:
        raise RuntimeError("Source freeze workbook changed during build")
    return {
        "output": str(output),
        "source_sha256": source_hash_before,
        "sita_rows": int(len(sita)),
        "bl8_rows": int(len(bl8)),
        "unified_rows": int(len(unified)),
        "method_counts": {str(k): int(v) for k, v in unified["method"].value_counts().to_dict().items()},
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--bl8-root", type=Path, default=DEFAULT_BL8_ROOT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    print(json.dumps(build(args.source, args.bl8_root, args.output), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
